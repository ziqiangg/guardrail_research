"""Independent verification of the Litmus apply (sheets 3m and 3n; NO sheet 3 columns; all other sheets vs the
previous commit). One-shot like the other verify_*_apply scripts.

Baseline = the workbook at git HEAD (the previous commit), not a frozen file under baselines/.
Writes drafts/litmus_urls.txt (3m Source URL cells and 3n URLs only; no R9 cells since there are no columns).
"""
import contextlib
import io
import os
import subprocess
import sys
import tempfile
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter

import build_two_level as B2
import build_inventory as BI
import build_litmus_inventory as BL
import build_litmus_eval as BE
import build_eval_sheet as ES
import inventory_sheet as IS
import products

XLSX = B2.XLSX
URLS = str(BI.P.DRAFTS / "litmus_urls.txt")
OLD_SHEETS = ["3. Guardrail Research Table", "3b. NeMo Rail Inventory", "3c. NeMo Evaluation Tooling",
              "3d. Llama Guard Inventory", "3e. GovTech Sentinel Inventory", "3f. Presidio Inventory",
              "3g. SDP Inventory", "3h. Model Armor Inventory", "3i. LionGuard Inventory",
              "3j. Purple Llama Inventory", "3k. CyberSecEval Eval Tooling", "3l. Cloak Inventory",
              "4. Candidate Comparison Groups"]
S3, S3M, S3N = OLD_SHEETS[0], BL.S3M, BE.S3N
NEW_SHEETS = [S3M, S3N]
RESULTS = []


def check(name, ok, extra=""):
    RESULTS.append(bool(ok))
    print(("PASS" if ok else "FAIL"), "-", name, extra)


plain = B2.plain_of


def runs_of(v):
    return [(b.text, bool(b.font.b)) for b in v] if isinstance(v, CellRichText) else [(plain(v), False)]


def _col(c):
    return None if c is None else (c.rgb if c.type == "rgb" else (c.type, c.theme if c.type == "theme" else c.indexed))


def style_of(c):
    b = c.border
    return (c.font.name, c.font.sz, c.font.b, c.font.i, _col(c.font.color), c.fill.fill_type, _col(c.fill.fgColor),
            c.alignment.wrap_text, c.alignment.vertical, c.alignment.horizontal,
            tuple((s.style, _col(s.color)) if s is not None else None for s in (b.left, b.right, b.top, b.bottom)))


def full_snap(w):
    cells = [(c.coordinate, plain(c.value), runs_of(c.value), style_of(c)) for row in w.iter_rows() for c in row]
    layout = (sorted(str(m) for m in w.merged_cells.ranges), w.freeze_panes, w.auto_filter.ref,
              {k: v.width for k, v in w.column_dimensions.items()},
              {k: (v.height, v.outline_level, v.hidden) for k, v in w.row_dimensions.items()},
              (w.max_row, w.max_column))
    return cells, layout


# ---------------------------------------------------------------- previous commit's workbook
prev = os.path.join(tempfile.mkdtemp(), "prev.xlsx")
with open(prev, "wb") as f:
    f.write(subprocess.run(["git", "show", 'HEAD:benchtest/AI Guardrails Research and Comparison.xlsx'],
                           capture_output=True, check=True, cwd=str(BI.P.ROOT)).stdout)
new = load_workbook(XLSX, rich_text=True)
old = load_workbook(prev, rich_text=True)
print("previous-commit sheet names:", old.sheetnames)
want_order = OLD_SHEETS[:12] + NEW_SHEETS + OLD_SHEETS[12:]
check("sheet order (3m right after 3l, 3n right after 3m, sheet 4 last)",
      new.sheetnames == list(BI.ORDER) == want_order, str(new.sheetnames[-4:]))
check("previous sheets all present, in order, plus only 3m and 3n",
      [n for n in new.sheetnames if n not in NEW_SHEETS] == old.sheetnames == OLD_SHEETS)

# ---------------------------------------------------------------- registry: no Table 3 columns
lt = products.PRODUCTS[9]
mds = [B2.parse_md(m) for m in B2.MDS]
check("registry: litmus entry has no header prefix and no two-level md; MDS has the 9 earlier products only; "
      "earlier 9 entries in order",
      lt["slug"] == "litmus" and lt["two_level_md"] is None and products.prefixes(lt) == ()
      and len(products.PRODUCTS) == 10 and len(products.MDS) == 9 and len(mds) == 9
      and [p["slug"] for p in products.PRODUCTS[:9]] == ["nemo", "llamaguard", "sentinel", "presidio", "sdp",
                                                         "modelarmor", "lionguard", "purplellama", "cloak"]
      and lt["inventory_sheet"] == S3M and lt["inventory_builder_module"] == "build_litmus_inventory")
check("registry: EXTRA_SHEETS has 3n after 3m", any(e["sheet"] == S3N and e["after"] == S3M
                                                    for e in products.EXTRA_SHEETS))

# ---------------------------------------------------------------- sheet 3 unchanged (no new columns)
ws, wo = new[S3], old[S3]
check("sheet 3: still 66 columns (A-BN), no new columns", ws.max_column == wo.max_column == 66,
      "%d vs %d" % (ws.max_column, wo.max_column))
sa, sal = full_snap(ws)
sb, sbl = full_snap(wo)
check("sheet 3: values, rich text runs, styles identical to previous commit", sa == sb,
      "(%d cells, %d diffs)" % (len(sa), sum(1 for x, y in zip(sa, sb) if x != y)))
check("sheet 3: merges, freeze, autofilter, widths, row heights/outline, extents identical", sal == sbl)

# ---------------------------------------------------------------- every other pre-existing sheet identical
for nm in OLD_SHEETS[1:]:
    a, al = full_snap(new[nm])
    b, bl = full_snap(old[nm])
    nd = sum(1 for x, y in zip(a, b) if x != y) + abs(len(a) - len(b))
    check(f"{nm}: values (incl. formulas), rich text, styles identical to previous commit", a == b,
          f"({len(a)} cells, {nd} diffs)")
    check(f"{nm}: merges, freeze, autofilter, widths, row heights/outline, extents identical", al == bl,
          "" if al == bl else str([(i, x, y) for i, (x, y) in enumerate(zip(al, bl)) if x != y])[:300])

# ---------------------------------------------------------------- 3m
w3m = new[S3M]
blocks = IS.parse_md(BL.CFG)
pos, prev_last = [], 2
for b in blocks:
    h, rows = IS.block_rows(w3m, b["hdr"][0], prev_last + 1)
    pos.append((h, rows))
    prev_last = rows[-1]
cnt = [len(p[1]) for p in pos]
check("3m block row counts 4/5/6 = 15", cnt == [4, 5, 6] and sum(cnt) == 15, f"{cnt}")
mm = tot = 0
for b, (h, rows) in zip(blocks, pos):
    got = [[w3m.cell(row=r, column=c).value for c in range(1, len(b["hdr"]) + 1)] for r in [h] + rows]
    want = [b["hdr"]] + b["rows"]
    tot += sum(len(x) for x in want)
    mm += sum(1 for g, w in zip(got, want) for x, y in zip(g, w) if x != y) + abs(len(got) - len(want))
check("3m every cell equals cleaned md text (%d cells)" % tot, mm == 0, f"({mm} mismatches)")
h0, r0 = pos[0]
check("3m autofilter = block (a) only",
      w3m.auto_filter.ref == f"A{h0}:{get_column_letter(len(blocks[0]['hdr']))}{r0[-1]}", w3m.auto_filter.ref)
check("3m Covered-by column in all three blocks; block widths 8/7/7",
      all(BL.COV in b["hdr"] for b in blocks) and [len(b["hdr"]) for b in blocks] == [8, 7, 7])
hdr3 = {ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)}
vals, amber_bad = [], 0
for b, (h, rows) in zip(blocks, pos):
    cc = b["hdr"].index(BL.COV) + 1
    for r in rows:
        v = w3m.cell(row=r, column=cc)
        vals.append(v.value)
        amber_bad += not (v.fill.fill_type == "solid" and str(v.fill.fgColor.rgb).endswith("FFF2CC"))
check("3m Covered-by: all 15 cells are exactly the inventory-only marker (no Table 3 header cited, so no "
      "header lookup applies); markers tuple = (inventory only,)",
      vals == [BL.INV_ONLY] * 15 and BL.CFG["markers"] == (BL.INV_ONLY,) and not any(v in hdr3 for v in vals))
check("3m amber fill on all 15 marker cells", amber_bad == 0)
ctx = dict(ws=w3m, ws3=ws, cfg=BL.CFG, blocks=blocks, pos=pos, hdr3=hdr3)
for fn in BL.CFG["validators"]:
    ok_, msgs = fn(ctx)
    check("3m validator " + fn.__name__, ok_, "| " + " ".join(msgs)[:300])
ok_p, counts, totals = IS.verify_panel(new, BL.CFG, w3m, ws, blocks, pos)
check("3m panel (zero Table 3 columns) equals spec: COUNTA per block 4/5/6, inventory-only total 15, no sheet-3 refs",
      ok_p and [v[0] for v in counts.values()] == [4, 5, 6] and totals["Rows marked inventory only"] == 15
      and not any("Guardrail Research Table" in str(c.value) for row in w3m.iter_rows() for c in row), str(totals))
mk = [c.coordinate for row in w3m.iter_rows() for c in row if isinstance(c.value, str) and ("**" in c.value or "`" in c.value)]
check("3m no ** or backticks", not mk, str(mk[:5]))
nonar = [c.coordinate for row in w3m.iter_rows() for c in row
         if type(c).__name__ != "MergedCell" and c.value is not None and c.font.name != "Arial"]
nx, bx = ES.xml_non_arial(XLSX, S3M)
check("3m all Arial (cells + raw XML <c> check)", not nonar and bx == 0, f"(openpyxl {len(nonar)}, raw XML {bx} of {nx})")
check("3m A1 bold 13; A2 italic grey; freeze A4",
      w3m["A1"].font.b and w3m["A1"].font.sz == 13 and w3m["A2"].font.i and w3m["A2"].font.color.rgb.endswith("808080")
      and w3m.freeze_panes == "A4")

# ---------------------------------------------------------------- 3n
buf = io.StringIO()
with contextlib.redirect_stdout(buf):
    try:
        mism_n = ES.verify(new, None, XLSX, BE.CFG)
        err = None
    except AssertionError as e:  # pragma: no cover
        mism_n, err = -1, e
out = buf.getvalue()
secs = ES.parse_md(BE.MD)
heads = "section headings: ['Overview', 'Tools', 'Datasets', 'Published results', 'Red-teaming', 'Engine coverage', " \
        "'Reuse for the test bench', 'Open questions'] OK"
check("3n parses to 8 sections in default order; sheet text equals md (0 mismatches, 0 row-count failures)",
      err is None and mism_n == 0 and len(secs) == 8 and "table row-count failures: 0" in out and "FAIL" not in out
      and heads in out,
      "(" + [l for l in out.splitlines() if l.startswith("plain-text")][0] + ")")
w3n = new[S3N]
nonar_n = [c.coordinate for row in w3n.iter_rows() for c in row
           if type(c).__name__ != "MergedCell" and c.value is not None and c.font.name != "Arial"]
nx, bx = ES.xml_non_arial(XLSX, S3N)
check("3n all Arial (cells + raw XML), title/note as configured, freeze A3",
      not nonar_n and bx == 0 and w3n["A1"].value == BE.TITLE and w3n["A2"].value == BE.NOTE
      and w3n.freeze_panes == "A3", f"(raw XML {bx} of {nx})")

# ---------------------------------------------------------------- XML hygiene
ssp, emp, norpr, n, wf = BI.hygiene(XLSX)
check("XML hygiene (all parts well-formed; worksheets + sharedStrings)", ssp == 0 and emp == 0 and norpr == 0 and wf,
      f"(unpreserved <t>={ssp}, <r> without rPr={norpr}, empty <t>={emp}, parts checked={n}, well-formed={wf})")

# ---------------------------------------------------------------- URLs (3m Source URL cells and 3n only)
urls, seen = [], set()


def skipped(u):
    p = urlparse(u)
    host = p.hostname or ""
    return ("{" in u or "}" in u or host == "form.gov.sg"
            or (host.startswith("litmus.") and host.endswith(".aiguardian.gov.sg"))
            or host.endswith("googleapis.com") or "/api/" in p.path)


def add(t):
    for u in BI.URL_RE.findall(t or ""):
        u = u.rstrip(".:,;)")
        if skipped(u):
            continue
        if u not in seen:
            seen.add(u)
            urls.append(u)


for b, (h, rows) in zip(blocks, pos):
    cc = b["hdr"].index("Source URL") + 1
    for r in rows:
        add(w3m.cell(row=r, column=cc).value)
n3m = len(urls)
for row in w3n.iter_rows():
    for c in row:
        if type(c).__name__ != "MergedCell":
            add(plain(c.value))
open(URLS, "w", encoding="utf-8", newline="\n").write("\n".join(urls) + "\n")
print(f"urls: {len(urls)} unique ({n3m} from 3m Source URL cells, {len(urls) - n3m} new from 3n) -> {URLS}")
check("URL list written, no skipped host/template entries", bool(urls) and not any(skipped(u) for u in urls),
      f"({len(urls)} URLs)")

print("\nSUMMARY: %d/%d checks passed" % (sum(RESULTS), len(RESULTS)))
sys.exit(0 if all(RESULTS) else 1)
