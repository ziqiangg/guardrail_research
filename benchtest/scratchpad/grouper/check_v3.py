# -*- coding: utf-8 -*-
"""Self-check for benchtest/drafts/groups_v3.md (independent parser; reads the workbook read-only).
Run from the repo root:  PYTHONIOENCODING=utf-8 python benchtest/scratchpad/grouper/check_v3.py
Exit code 1 if any hard check fails."""
import os, re, sys

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
MD = os.path.join(ROOT, "benchtest", "drafts", "groups_v3.md")
XLSX = os.path.join(ROOT, "benchtest", "AI Guardrails Research and Comparison.xlsx")
MARKER = "• Single-product: no comparator among the ten products yet"
MAXW = 12
errors, warns = [], []


def err(m):
    errors.append(m)


def warn(m):
    warns.append(m)


def split_row(line):
    s = line.strip()
    assert s.startswith("|") and s.endswith("|"), s[:40]
    return [c.strip() for c in s[1:-1].split(" | ")] if " | " in s else [c.strip() for c in s[1:-1].split("|")]


def colnum(L):
    n = 0
    for ch in L:
        n = n * 26 + ord(ch) - 64
    return n


# ---- workbook headers (read-only) -------------------------------------------------------------------------
import openpyxl
wb = openpyxl.load_workbook(XLSX, read_only=True)
ws = wb["3. Guardrail Research Table"]
row3 = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))
HDR = {}
from openpyxl.utils import get_column_letter
for c in range(5, len(row3) + 1):
    v = row3[c - 1]
    if isinstance(v, str) and v.strip():
        HDR[get_column_letter(c)] = v.strip()
wb.close()
LAST = max(HDR, key=colnum)
print(f"sheet 3 function columns: {len(HDR)} (E to {LAST})")
if colnum(LAST) - 4 != len(HDR):
    err("gaps in sheet 3 header row")

# header uniqueness (COUNTIF substring safety, as build_groups_sheet.fn_headers)
for h in HDR.values():
    if any(ch in h for ch in "*?~"):
        err(f"wildcard char in header {h}")
    for o in HDR.values():
        if h != o and h.lower() in o.lower():
            err(f"header is a substring of another: {h}")

text = open(MD, encoding="utf-8").read()
lines = text.splitlines()

# ---- section A table --------------------------------------------------------------------------------------
i = next(k for k, l in enumerate(lines) if l.startswith("## A."))
j = next(k for k in range(i + 1, len(lines)) if lines[k].startswith("|"))
tbl = []
while j < len(lines) and lines[j].startswith("|"):
    tbl.append(split_row(lines[j]))
    j += 1
hdr, rows = tbl[0], tbl[2:]
EXPECT_HDR = ["Candidate group", "Guardrail functions included", "Common test inputs", "Ground truth", "Outputs to capture",
              "Common metrics", "Minimum architecture", "Material differences or limitations"]
if hdr != EXPECT_HDR:
    err(f"table header differs from the 8 template columns: {hdr}")
if not all(set(x) <= set("-: ") for x in tbl[1]):
    err("separator row malformed")
print(f"table rows: {len(rows)}")

in_group = {}
nwords_max = 0
bullets = 0
nsingle = 0
seen_single = False
group_funcs = {}
for k, r in enumerate(rows):
    if len(r) != 8:
        err(f"row {k + 1}: {len(r)} cells")
        continue
    name = r[0]
    cid = name.split()[0]
    if cid != f"C{k + 1}":
        err(f"row {k + 1}: id {cid} out of sequence")
    # functions
    letters = []
    for part in r[1].split("; "):
        m = re.match(r"([A-Z]{1,2}): (.*)$", part)
        if not m:
            err(f"{cid}: bad function entry {part[:40]}")
            continue
        L, h = m.groups()
        if L not in HDR:
            err(f"{cid}: unknown column {L}")
            continue
        if not h.startswith(HDR[L]):
            err(f"{cid}: header text of {L} differs from sheet 3: {h[:60]}")
        letters.append(L)
        in_group.setdefault(L, []).append(cid)
    if len(set(letters)) != len(letters):
        err(f"{cid}: duplicate function")
    group_funcs[cid] = letters
    single = r[7].split("<br>")[0].strip() == MARKER
    if r[7].lstrip().startswith("• Single-product:") and not single:
        err(f"{cid}: marker bullet is not the exact v2 marker")
    if single:
        nsingle += 1
        seen_single = True
        if len(letters) < 1:
            err(f"{cid}: no functions")
    elif seen_single:
        err(f"{cid}: multi-product row after a single-product row (band order)")
    PROD = {"LlamaFirewall": "Purple Llama", "Code Shield": "Purple Llama", "Prompt Guard 2": "Purple Llama"}
    prods = {PROD.get(HDR[L].split(":")[0], HDR[L].split(":")[0]) for L in letters}
    if not single and len(prods) < 2:
        err(f"{cid}: marked multi-product but spans only {prods}")
    if single and len(prods) != 1:
        err(f"{cid}: single-product row spans {prods}")
    if single:
        prefixes = {HDR[L].split(":")[0] for L in letters}
        if len(prefixes) > 1 and not all(p in ("LlamaFirewall", "Code Shield", "Prompt Guard 2") for p in prefixes):
            err(f"{cid}: single-product row spans prefixes {prefixes}")
    # text cells
    for c in range(2, 8):
        parts = [x.strip() for x in r[c].split("<br>")]
        if "**" in r[c] or "`" in r[c]:
            err(f"{cid} col {c + 1}: ** or backtick")
        if not parts or not parts[-1].startswith("Refs: ") or len(parts[-1]) <= 6:
            err(f"{cid} col {c + 1}: no final Refs line")
        for n, p in enumerate(parts):
            if p.startswith("Refs: "):
                if n != len(parts) - 1:
                    err(f"{cid} col {c + 1}: Refs not last")
                continue
            if not p.startswith("• ") or len(p) <= 2:
                err(f"{cid} col {c + 1}: bullet missing marker: {p[:40]}")
                continue
            w = len(p[2:].split())
            bullets += 1
            nwords_max = max(nwords_max, w)
            if w > MAXW:
                err(f"{cid} col {c + 1}: {w} words: {p[:70]}")
        if c == 7:
            pass
    # cross references to groups
    for m in re.finditer(r"\bC(\d+)\b", " ".join(r[2:])):
        if not 1 <= int(m.group(1)) <= len(rows):
            err(f"{cid}: reference to missing group C{m.group(1)}")
    # refs letters
    for c in range(2, 8):
        last = r[c].split("<br>")[-1]
        for m in re.finditer(r"\b([A-Z]{1,2}) R([0-9])\b", last):
            if m.group(1) not in HDR:
                err(f"{cid}: Refs cite unknown column {m.group(1)}")
            elif m.group(2) == "0":
                err(f"{cid}: bad row ref")
            elif m.group(1) not in letters:
                warn(f"{cid} col {c + 1}: Refs cite {m.group(1)} R{m.group(2)}, not a member")
print(f"bullets checked: {bullets}; longest bullet: {nwords_max} words (limit {MAXW})")
print(f"multi-product groups: {len(rows) - nsingle}; single-product groups: {nsingle}")

# ---- coverage ---------------------------------------------------------------------------------------------
missing = [L for L in HDR if L not in in_group]
print(f"columns in no group: {missing if missing else 'none'}")
if missing:
    err(f"columns not covered: {missing}")
multi_cols = [L for L, g in in_group.items() if len(g) > 1]
print(f"columns in more than one group: {len(multi_cols)}")

# ---- accounting table matches the group table -------------------------------------------------------------
a = next(k for k, l in enumerate(lines) if l.startswith("## C."))
acct = {}
for l in lines[a:]:
    if l.startswith("## D."):
        break
    if l.startswith("| ") and not l.startswith("| Sheet 3 column") and not l.startswith("|---"):
        cells = split_row(l)
        if len(cells) == 3:
            acct[cells[0]] = (cells[1], [x.strip() for x in cells[2].split(",") if x.strip()])
if set(acct) != set(HDR):
    err(f"accounting table columns differ: missing {set(HDR) - set(acct)}, extra {set(acct) - set(HDR)}")
for L, (h, gs) in acct.items():
    if L in HDR and h != HDR[L]:
        err(f"accounting header differs for {L}")
    if sorted(gs, key=lambda x: int(x[1:])) != in_group.get(L, []) and sorted(gs) != sorted(in_group.get(L, [])):
        err(f"accounting groups differ for {L}: {gs} vs {in_group.get(L)}")
print(f"accounting rows: {len(acct)}")

# ---- rationale ---------------------------------------------------------------------------------------------
b = next(k for k, l in enumerate(lines) if l.startswith("## B."))
blocks = {}
cur = None
for l in lines[b:a]:
    m = re.match(r"### (C\d+) ", l)
    if m:
        cur = m.group(1)
        blocks[cur] = []
    elif cur:
        blocks[cur].append(l)
for k in range(1, len(rows) + 1):
    cid = f"C{k}"
    body = "\n".join(blocks.get(cid, []))
    if not body:
        err(f"no rationale for {cid}")
        continue
    ticks = body.count("✓") + body.count("✗")
    if ticks < 4:
        err(f"{cid}: rationale has fewer than four tick or cross marks")
    if "R" not in body:
        err(f"{cid}: rationale lacks sheet 3 references")
print(f"rationale blocks: {len(blocks)}")

# ---- sections and characters -----------------------------------------------------------------------------
for s in ("## D. Changes against v2",):
    if s not in text:
        err(f"missing section {s}")
allowed = set("•—–✓✗")
bad = sorted({ch for ch in text if ord(ch) > 127 and ch not in allowed})
if bad:
    err(f"unexpected non-ASCII characters: {bad}")
if re.search(r"[\U0001F000-\U0001FFFF]", text):
    err("emoji present")
for w in ("anonymization", "color ", "organization", "behavior", "licence_"):
    pass

print("WARNINGS (refs to non-members, informational):", len(warns))
for w in warns[:60]:
    print("  warn:", w)
print("ERRORS:", len(errors))
for e in errors:
    print("  ERROR:", e)
sys.exit(1 if errors else 0)
