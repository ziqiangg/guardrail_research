# Legacy: verified the Llama Guard change against v4; superseded by verify_sentinel_apply.py / verify_groups_apply.py;
# not runnable against the current workbook (it expects the old 26-column workbook and old hard-coded paths).
"""Independent verification of the Llama Guard apply (sheet 3 V-Z, sheet 3d, untouched sheets vs .v4)."""
import re, sys
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import get_column_letter

import build_two_level as B2
import build_inventory as BI
from fill_nemo_columns import LABEL_RE

DIR = B2.DIR
XLSX = B2.XLSX
V4 = DIR + r"\AI Guardrails Research and Comparison.v4.xlsx"
URLS = DIR + r"\drafts\lg_urls.txt"
S3, S3B, S3C, S3D, S4 = B2.S3, BI.S3B, BI.S3C, BI.S3D, B2.S4
RESULTS = []


def check(name, ok, extra=""):
    RESULTS.append(bool(ok))
    print(("PASS" if ok else "FAIL"), "-", name, extra)


plain = B2.plain_of


def runs_of(v):
    return [(b.text, bool(b.font.b)) for b in v] if isinstance(v, CellRichText) else [(plain(v), False)]


def snap(w, fonts=True):
    out = []
    for row in w.iter_rows():
        for c in row:
            if fonts:
                out.append((c.coordinate, plain(c.value), c.font.name, c.font.sz, c.font.b, c.font.i,
                            c.fill.fill_type, c.fill.fgColor.rgb))
            else:
                out.append((c.coordinate, c.value))
    return out, sorted(str(m) for m in w.merged_cells.ranges)


new = load_workbook(XLSX, rich_text=True)
old = load_workbook(V4, rich_text=True)
print("v4 sheet names:", old.sheetnames)
want_order = [S3, S3B, S3C, S3D, S4]
check("sheet order", new.sheetnames == want_order, str(new.sheetnames))
check("v4 names for 3/3b/3c/4 unchanged", old.sheetnames == [S3, S3B, S3C, S4])

# ---------------------------------------------------------------- sheet 3 structure
ws, wo = new[S3], old[S3]
check("sheet 3 max_column 26 (A-Z)", ws.max_column == 26, str(ws.max_column))
check("sheet 3 max_row 21", ws.max_row == 21, str(ws.max_row))
check("27 merges", len(ws.merged_cells.ranges) == 27 and
      sorted(map(str, ws.merged_cells.ranges)) == sorted(map(str, wo.merged_cells.ranges)))
lv = [r for r in range(1, 23) if ws.row_dimensions[r].outline_level == 1]
check("outline levels (rows 5,7,..,21 level 1; others 0)", lv == list(range(5, 22, 2)), str(lv))
check("freeze panes E4 as v4", ws.freeze_panes == wo.freeze_panes == "E4")
check("summaryBelow False", ws.sheet_properties.outlinePr.summaryBelow is False)
widths = [ws.column_dimensions[get_column_letter(i)].width for i in range(1, 27)]
check("column widths 6/28/32/10 then 48", widths == [6, 28, 32, 10] + [48] * 22, str(widths))

md_lg = B2.parse_md(B2.MDS[1])
lg_hdrs = [c["header"] for c in md_lg]
got_hdr = [ws.cell(row=3, column=c).value for c in range(22, 27)]
check("V3..Z3 == 5 LG headers in order", got_hdr == lg_hdrs and len(lg_hdrs) == 5 and
      all(h.startswith("Llama Guard:") for h in got_hdr))
nemo = [ws.cell(row=3, column=c).value for c in range(6, 22)]
check("F3..U3 are the 16 NeMo headers", len(nemo) == 16 and all(h.startswith("NeMo Guardrails:") for h in nemo))

# ---------------------------------------------------------------- NeMo A-U vs v4
diff = 0
for r in range(3, 22):
    for c in range(1, 22):
        diff += plain(ws.cell(row=r, column=c).value) != plain(wo.cell(row=r, column=c).value)
check("A-U rows 3-21 plain text vs v4: 0 diffs", diff == 0, f"({diff} diffs)")
rdiff = 0
for r in range(4, 22):
    for c in range(6, 22):
        rdiff += runs_of(ws.cell(row=r, column=c).value) != runs_of(wo.cell(row=r, column=c).value)
check("F-U rows 4-21 rich-text runs (text+bold), all cells, vs v4: 0 diffs", rdiff == 0, f"({rdiff} diffs)")
sdiff = 0
for r in range(3, 22):
    for c in range(1, 22):
        a, b = ws.cell(row=r, column=c), wo.cell(row=r, column=c)
        sdiff += (a.font.name, a.font.sz, a.font.b, a.fill.fill_type, a.fill.fgColor.rgb, a.alignment.wrap_text) != \
                 (b.font.name, b.font.sz, b.font.b, b.fill.fill_type, b.fill.fgColor.rgb, b.alignment.wrap_text)
check("A-U cell styles (font/fill/wrap) vs v4: 0 diffs", sdiff == 0, f"({sdiff} diffs)")

# ---------------------------------------------------------------- LG V-Z vs md
mism = 0
for n in range(1, 10):
    sr, dr = 4 + 2 * (n - 1), 5 + 2 * (n - 1)
    for ci, col in enumerate(md_lg):
        cc = 22 + ci
        d = col["R"][n]
        _, sp = B2.rich(d["summary"], 10, "v")
        _, dp = B2.rich("\n".join(d["detail"]), 9, "v")
        mism += plain(ws.cell(row=sr, column=cc).value) != sp
        mism += plain(ws.cell(row=dr, column=cc).value) != dp
check("V-Z plain text vs lg_two_level.md parse: 0 mismatches (90 cells)", mism == 0, f"({mism})")
fails = 0
for n in range(1, 10):
    sr = 4 + 2 * (n - 1)
    for cc in range(22, 27):
        c = ws.cell(row=sr, column=cc)
        blocks = list(c.value) if isinstance(c.value, CellRichText) else [c.value]
        p = plain(c.value)
        probs = []
        if n <= 8 and not (isinstance(blocks[0], TextBlock) and blocks[0].font.b):
            probs.append("first-not-bold")
        if n <= 7 and not any(isinstance(b, TextBlock) and b.font.b and LABEL_RE.search(b.text) for b in blocks):
            probs.append("no-bold-label")
        if any(s in p for s in ("**", "`", "_", "$")):
            probs.append("marker")
        if len(p.split()) > 60:
            probs.append("words=%d" % len(p.split()))
        if "\n" in p:
            probs.append("newline")
        if probs:
            fails += 1
            print("   ", c.coordinate, probs)
check("V-Z Summary style checks (45 cells)", fails == 0, f"({fails} failing)")
nonar = [ws.cell(row=r, column=c).coordinate for r in range(3, 22) for c in range(22, 27)
         if ws.cell(row=r, column=c).font.name != "Arial"]
check("V-Z non-Arial cells: 0", not nonar, str(nonar[:5]))

# ---------------------------------------------------------------- 3b vs v4
a, am = snap(new[S3B], fonts=False)
b, bm = snap(old[S3B], fonts=False)
bd = [x for x, y in zip(a, b) if x != y]
check("3b every cell value (incl. formula strings) vs v4: 0 diffs", not bd and len(a) == len(b) and am == bm,
      f"({len(bd)} diffs, {len(a)} vs {len(b)} cells, merges equal={am == bm})")
a, _ = snap(new[S3B])
b, _ = snap(old[S3B])
check("3b fonts/fills vs v4: 0 diffs", a == b)
w3b = new[S3B]
labels = [(c.row, c.value) for row in w3b.iter_rows(min_col=1, max_col=1) for c in row
          if isinstance(c.value, str) and c.value.startswith("='" + S3 + "'!")]
cols = [re.search(r"!([A-Z]+)3$", v).group(1) for _, v in labels]
exp_cols = [get_column_letter(i) for i in range(6, 22)]
check("3b panel C lists exactly 16 label formulas ='3...'!F3..U3", cols == exp_cols, f"({len(labels)}: {cols[0]}..{cols[-1]})")
check("3b panel labels resolve to the 16 NeMo headers", [ws[c + "3"].value for c in cols] == nemo)
fv = [c.value for row in w3b.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
fo = [c.value for row in old[S3B].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
check("3b formula strings identical to v4", fv == fo, f"({len(fv)} formulas)")

# ---------------------------------------------------------------- 3c, 4 vs v4
for nm in (S3C, S4):
    a, am = snap(new[nm])
    b, bm = snap(old[nm])
    check(f"{nm}: values, fonts, fills, merges identical to v4", a == b and am == bm,
          f"({sum(1 for x, y in zip(a, b) if x != y)} diffs)")
    check(f"{nm}: freeze panes / widths same as v4", new[nm].freeze_panes == old[nm].freeze_panes and
          {k: v.width for k, v in new[nm].column_dimensions.items()} ==
          {k: v.width for k, v in old[nm].column_dimensions.items()})

# ---------------------------------------------------------------- XML hygiene
ssp, emp, norpr, n, wf = BI.hygiene(XLSX)
check("XML hygiene (all parts well-formed; ws+sharedStrings)", ssp == 0 and emp == 0 and norpr == 0 and wf,
      f"(unpreserved <t>={ssp}, <r> without rPr={norpr}, empty <t>={emp}, parts checked={n}, well-formed={wf})")

# ---------------------------------------------------------------- URLs
urls, seen = [], set()


def add(t):
    for u in BI.URL_RE.findall(t or ""):
        u = u.rstrip(".:")
        if u not in seen:
            seen.add(u)
            urls.append(u)


for c in range(22, 27):
    add(plain(ws.cell(row=21, column=c).value))
n9 = len(urls)
for row in new[S3D].iter_rows():
    for c in row:
        add(c.value if isinstance(c.value, str) else None)
open(URLS, "w", encoding="utf-8", newline="\n").write("\n".join(urls) + "\n")
print(f"urls: {len(urls)} unique ({n9} from V21:Z21, {len(urls) - n9} new from 3d) -> {URLS}")

print("\nSUMMARY: %d/%d checks passed" % (sum(RESULTS), len(RESULTS)))
sys.exit(0 if all(RESULTS) else 1)
