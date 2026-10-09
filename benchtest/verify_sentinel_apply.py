"""Independent verification of the GovTech Sentinel apply (sheet 3 AA-AG, sheet 3e, untouched sheets vs .v5)."""
import re, sys
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import get_column_letter

import build_two_level as B2
import build_inventory as BI
import build_sentinel_inventory as BS
import inventory_sheet as IS
from fill_nemo_columns import LABEL_RE

DIR = B2.DIR
XLSX = B2.XLSX
V5 = BI.BAK5
URLS = DIR + r"\drafts\sentinel_urls.txt"
S3, S3B, S3C, S3D, S3E, S4 = B2.S3, BI.S3B, BI.S3C, BI.S3D, BI.S3E, B2.S4
RESULTS = []


def check(name, ok, extra=""):
    RESULTS.append(bool(ok))
    print(("PASS" if ok else "FAIL"), "-", name, extra)


plain = B2.plain_of


def runs_of(v):
    return [(b.text, bool(b.font.b)) for b in v] if isinstance(v, CellRichText) else [(plain(v), False)]


def _col(c):
    return None if c is None else (c.rgb if c.type == "rgb" else (c.type, c.theme if c.type == "theme" else c.indexed))


def style_of(c):
    b = c.border
    return (c.font.name, c.font.sz, c.font.b, c.font.i, _col(c.font.color), c.fill.fill_type, _col(c.fill.fgColor),
            c.alignment.wrap_text, c.alignment.vertical, c.alignment.horizontal,
            tuple((s.style, _col(s.color)) if s is not None else None for s in (b.left, b.right, b.top, b.bottom)))


def full_snap(w, max_row=None):
    """Everything observable on a sheet: values (incl. formula strings), styles, merges, layout.
    max_row limits the snapshot to the original rows (excludes a later-appended panel and the extents)."""
    top = max_row or w.max_row
    cells = [(c.coordinate, plain(c.value), style_of(c)) for row in w.iter_rows(max_row=top) for c in row]
    layout = (sorted(str(m) for m in w.merged_cells.ranges if m.min_row <= top), w.freeze_panes, w.auto_filter.ref,
              {k: v.width for k, v in w.column_dimensions.items()},
              {k: (v.height, v.outline_level, v.hidden) for k, v in w.row_dimensions.items() if k <= top},
              (top, w.max_column) if max_row else (w.max_row, w.max_column))
    return cells, layout


new = load_workbook(XLSX, rich_text=True)
old = load_workbook(V5, rich_text=True)
print("v5 sheet names:", old.sheetnames)
want_order = [S3, S3B, S3C, S3D, S3E, S4]
check("sheet order", new.sheetnames == want_order, str(new.sheetnames))
check("v5 sheet names unchanged in the new workbook", [n for n in new.sheetnames if n != S3E] == old.sheetnames
      == [S3, S3B, S3C, S3D, S4])

# ---------------------------------------------------------------- derive ranges from the drafts
md_nemo = B2.parse_md(B2.MDS[0])
md_lg = B2.parse_md(B2.MDS[1])
md_sn = B2.parse_md(B2.MDS[2])
n_nemo, n_lg, n_sn = len(md_nemo), len(md_lg), len(md_sn)
NEMO0, LG0 = 6, 6 + n_nemo
SN0 = LG0 + n_lg
LAST = SN0 + n_sn - 1
print("derived ranges: NeMo %s-%s (%d), LG %s-%s (%d), Sentinel %s-%s (%d), last col %d" % (
    get_column_letter(NEMO0), get_column_letter(LG0 - 1), n_nemo, get_column_letter(LG0),
    get_column_letter(SN0 - 1), n_lg, get_column_letter(SN0), get_column_letter(LAST), n_sn, LAST))

# ---------------------------------------------------------------- sheet 3 structure
ws, wo = new[S3], old[S3]
check("sheet 3: %d columns (A-%s)" % (LAST, get_column_letter(LAST)), ws.max_column == LAST == 33, str(ws.max_column))
check("sheet 3: max_row 21", ws.max_row == 21, str(ws.max_row))
check("27 merges equal to v5", len(ws.merged_cells.ranges) == 27 and
      sorted(map(str, ws.merged_cells.ranges)) == sorted(map(str, wo.merged_cells.ranges)))
lv = [r for r in range(1, 23) if ws.row_dimensions[r].outline_level == 1]
check("outline levels (rows 5,7,..,21 level 1; others 0)", lv == list(range(5, 22, 2)), str(lv))
check("freeze panes E4 (as v5)", ws.freeze_panes == wo.freeze_panes == "E4")
check("summaryBelow False", ws.sheet_properties.outlinePr.summaryBelow is False)
widths = [ws.column_dimensions[get_column_letter(i)].width for i in range(1, LAST + 1)]
check("column widths 6/28/32/10 then 48", widths == [6, 28, 32, 10] + [48] * (LAST - 4), str(widths[:6]) + "...")
hdr = lambda a, b: [ws.cell(row=3, column=c).value for c in range(a, b)]
check("F3..U3 are the %d NeMo headers (== md order)" % n_nemo,
      hdr(NEMO0, LG0) == [c["header"] for c in md_nemo] and n_nemo == 16
      and all(h.startswith("NeMo Guardrails:") for h in hdr(NEMO0, LG0)))
check("V3..Z3 are the %d Llama Guard headers" % n_lg,
      hdr(LG0, SN0) == [c["header"] for c in md_lg] and n_lg == 5 and all(h.startswith("Llama Guard:") for h in hdr(LG0, SN0)))
sn_hdrs = [c["header"] for c in md_sn]
check("AA3..AG3 are the 7 Sentinel headers in md order",
      hdr(SN0, LAST + 1) == sn_hdrs and n_sn == 7 and all(h.startswith("GovTech Sentinel:") for h in sn_hdrs),
      str([h[17:50] for h in sn_hdrs]))

# ---------------------------------------------------------------- A-Z vs v5
diff = sum(plain(ws.cell(row=r, column=c).value) != plain(wo.cell(row=r, column=c).value)
           for r in range(3, 22) for c in range(1, SN0))
check("A-Z rows 3-21 plain text vs v5: 0 diffs", diff == 0, f"({diff} diffs)")
rdiff = sum(runs_of(ws.cell(row=r, column=c).value) != runs_of(wo.cell(row=r, column=c).value)
            for r in range(4, 22) for c in range(6, SN0))
check("F-Z rows 4-21 rich-text runs (text+bold), all cells, vs v5: 0 diffs", rdiff == 0, f"({rdiff} diffs)")
sdiff = sum(style_of(ws.cell(row=r, column=c)) != style_of(wo.cell(row=r, column=c))
            for r in range(1, 22) for c in range(1, SN0))
check("A-Z rows 1-21 cell styles (font/fill/alignment/border) vs v5: 0 diffs", sdiff == 0, f"({sdiff} diffs)")

# ---------------------------------------------------------------- Sentinel AA-AG vs md
mism = ncells = 0
for n in range(1, 10):
    sr, dr = 4 + 2 * (n - 1), 5 + 2 * (n - 1)
    for ci, col in enumerate(md_sn):
        cc = SN0 + ci
        d = col["R"][n]
        _, sp = B2.rich(d["summary"], 10, "v")
        _, dp = B2.rich("\n".join(d["detail"]), 9, "v")
        mism += plain(ws.cell(row=sr, column=cc).value) != sp
        mism += plain(ws.cell(row=dr, column=cc).value) != dp
        ncells += 2
check("AA-AG plain text vs sentinel_two_level.md parse: 0 mismatches (%d cells)" % ncells,
      mism == 0 and ncells == 126, f"({mism})")
fails = nsum = 0
for n in range(1, 10):
    sr = 4 + 2 * (n - 1)
    for cc in range(SN0, LAST + 1):
        nsum += 1
        c = ws.cell(row=sr, column=cc)
        blocks = list(c.value) if isinstance(c.value, CellRichText) else [c.value]
        p = plain(c.value)
        probs = []
        if n <= 8 and not (isinstance(blocks[0], TextBlock) and blocks[0].font.b):
            probs.append("first-not-bold")
        if n <= 7 and not any(isinstance(b, TextBlock) and b.font.b and LABEL_RE.search(b.text) for b in blocks):
            probs.append("no-bold-label")
        if any(s in p for s in ("**", "`", "_", "$")):
            probs.append("marker")
        if len(p.split()) > 60:
            probs.append("words=%d" % len(p.split()))
        if "\n" in p:
            probs.append("newline")
        if probs:
            fails += 1
            print("   ", c.coordinate, probs)
check("AA-AG Summary style checks (%d cells)" % nsum, fails == 0 and nsum == 63, f"({fails} failing)")
bad_md = [ws.cell(row=r, column=c).coordinate for r in range(4, 22) for c in range(SN0, LAST + 1)
          if any(s in plain(ws.cell(row=r, column=c).value) for s in ("**", "`"))]
check("AA-AG no ** or backticks anywhere (rows 4-21)", not bad_md, str(bad_md[:5]))
nonar = [ws.cell(row=r, column=c).coordinate for r in range(3, 22) for c in range(SN0, LAST + 1)
         if ws.cell(row=r, column=c).font.name != "Arial"]
rich_nonar = sum(1 for r in range(4, 22) for c in range(SN0, LAST + 1)
                 for b in (ws.cell(row=r, column=c).value if isinstance(ws.cell(row=r, column=c).value, CellRichText) else [])
                 if isinstance(b, TextBlock) and b.font.rFont != "Arial")
check("AA-AG non-Arial cells / rich runs: 0", not nonar and rich_nonar == 0, f"({nonar[:5]}, rich runs {rich_nonar})")

# ---------------------------------------------------------------- 3b, 3c, 3d, 4 vs v5
# Sheet 4 is excluded (filled from groups.md after v5; verified by verify_groups_apply.py). 3d gained a coverage
# panel below its last block after v5, so it is compared on its original rows only.
for nm in (S3B, S3C, S3D):
    top = old[nm].max_row if nm == S3D else None
    a, al = full_snap(new[nm], top)
    b, bl = full_snap(old[nm], top)
    nd = sum(1 for x, y in zip(a, b) if x != y) + abs(len(a) - len(b))
    check(f"{nm}: cell values (incl. formula strings), fonts, fills, borders, alignment identical to v5",
          a == b, f"({len(a)} cells, {nd} diffs)")
    check(f"{nm}: merges, freeze, autofilter, widths, row heights/outline, extents identical to v5", al == bl,
          "" if al == bl else str([(i, x, y) for i, (x, y) in enumerate(zip(al, bl)) if x != y])[:300])
w3b = new[S3B]
labels = [(c.row, c.value) for row in w3b.iter_rows(min_col=1, max_col=1) for c in row
          if isinstance(c.value, str) and c.value.startswith("='" + S3 + "'!")]
cols = [re.search(r"!([A-Z]+)3$", v).group(1) for _, v in labels]
check("3b panel C lists exactly 16 label formulas ='3...'!F3..U3",
      cols == [get_column_letter(i) for i in range(NEMO0, LG0)], f"({len(labels)}: {cols[0]}..{cols[-1]})")
check("3b panel labels resolve to the 16 NeMo headers", [ws[c + "3"].value for c in cols] == hdr(NEMO0, LG0))

# ---------------------------------------------------------------- 3e
w3e = new[S3E]
blocks = IS.parse_md(BS.CFG)
pos = [IS.block_rows(w3e, b["hdr"][0]) for b in blocks]
check("3e block row counts 9/24/12/9", [len(p[1]) for p in pos] == [9, 24, 12, 9], str([len(p[1]) for p in pos]))
mm = tot = 0
for b, (h, rows) in zip(blocks, pos):
    got = [[w3e.cell(row=r, column=c).value for c in range(1, len(b["hdr"]) + 1)] for r in [h] + rows]
    want = [b["hdr"]] + b["rows"]
    tot += sum(len(x) for x in want)
    mm += sum(1 for g, w in zip(got, want) for x, y in zip(g, w) if x != y) + abs(len(got) - len(want))
check("3e every cell equals cleaned md text (%d cells)" % tot, mm == 0, f"({mm} mismatches)")
h0, r0 = pos[0]
check("3e autofilter = block (a) only", w3e.auto_filter.ref == f"A{h0}:P{r0[-1]}" and h0 == 3, w3e.auto_filter.ref)
hdr3 = {ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)}
badc, ncov, amber_bad = [], 0, 0
for bi in (0, 1):
    b, (h, rows) = blocks[bi], pos[bi]
    cc = b["hdr"].index("Covered by Table 3 column") + 1
    for r in rows:
        v = w3e.cell(row=r, column=cc)
        mk = v.value in (BS.LEGACY, BS.PLANNED)
        amber_bad += mk != (v.fill.fill_type == "solid" and str(v.fill.fgColor.rgb).endswith("FFF2CC"))
        for part in str(v.value).split(";"):
            part = part.strip()
            ncov += 1
            if part not in hdr3 and part not in (BS.LEGACY, BS.PLANNED):
                badc.append((r, part))
check("3e Covered-by values valid (real sheet-3 header or marker) in blocks (a),(b)", not badc,
      f"({ncov} values; bad {badc[:3]})")
check("3e amber fill exactly on marker cells", amber_bad == 0)
check("3e blocks (c),(d) have no Covered-by column",
      all("Covered by Table 3 column" not in blocks[i]["hdr"] for i in (2, 3)))
ctx = dict(ws=w3e, ws3=ws, cfg=BS.CFG, blocks=blocks, pos=pos, hdr3=hdr3)
for fn in BS.CFG["validators"]:
    ok_, msgs = fn(ctx)
    check("3e validator " + fn.__name__, ok_, "| " + " ".join(msgs)[:300])
mk = [c.coordinate for row in w3e.iter_rows() for c in row if isinstance(c.value, str) and ("**" in c.value or "`" in c.value)]
check("3e no ** or backticks", not mk, str(mk[:5]))
nonar = [c.coordinate for row in w3e.iter_rows() for c in row
         if type(c).__name__ != "MergedCell" and c.value is not None and c.font.name != "Arial"]
from build_eval_sheet import xml_non_arial
nx, bx = xml_non_arial(XLSX, S3E)
check("3e all Arial (cells + raw XML <c> check)", not nonar and bx == 0, f"(openpyxl {len(nonar)}, raw XML {bx} of {nx})")
nb = [r for r in range(1, w3e.max_row + 1) if w3e.cell(row=r, column=1).font.b]
check("3e A1 bold 13; A2 italic grey; bold titles above blocks 2-4; freeze A4",
      w3e["A1"].font.b and w3e["A1"].font.sz == 13 and w3e["A2"].font.i and w3e["A2"].font.color.rgb.endswith("808080")
      and [w3e.cell(row=p[0] - (2 if b["intro"] else 1), column=1).value for p, b in zip(pos[1:], blocks[1:])]
      == [b["title"] for b in blocks[1:]]
      and all(w3e.cell(row=p[0] - (2 if b["intro"] else 1), column=1).font.b for p, b in zip(pos[1:], blocks[1:]))
      and w3e.freeze_panes == "A4")

# ---------------------------------------------------------------- XML hygiene
ssp, emp, norpr, n, wf = BI.hygiene(XLSX)
check("XML hygiene (all parts well-formed; worksheets + sharedStrings)", ssp == 0 and emp == 0 and norpr == 0 and wf,
      f"(unpreserved <t>={ssp}, <r> without rPr={norpr}, empty <t>={emp}, parts checked={n}, well-formed={wf})")

# ---------------------------------------------------------------- URLs
urls, seen = [], set()


def add(t):
    for u in BI.URL_RE.findall(t or ""):
        u = u.rstrip(".:")
        if u not in seen:
            seen.add(u)
            urls.append(u)


for c in range(SN0, LAST + 1):
    add(plain(ws.cell(row=21, column=c).value))
n9 = len(urls)
for row in w3e.iter_rows():
    for c in row:
        add(c.value if isinstance(c.value, str) else None)
open(URLS, "w", encoding="utf-8", newline="\n").write("\n".join(urls) + "\n")
print(f"urls: {len(urls)} unique ({n9} from AA21:AG21, {len(urls) - n9} new from 3e) -> {URLS}")

print("\nSUMMARY: %d/%d checks passed" % (sum(RESULTS), len(RESULTS)))
sys.exit(0 if all(RESULTS) else 1)
