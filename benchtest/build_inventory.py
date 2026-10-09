"""Sheet '3b. NeMo Rail Inventory' from drafts/inventory.md. Called by build_two_level.main()."""
import json, os, re, shutil, subprocess, sys, time, zipfile
from collections import Counter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DIR = r"C:\Users\cys-c\Desktop\gzqr\benchtest"
INV = DIR + r"\drafts\inventory.md"
URLS_OUT = DIR + r"\drafts\v2_urls.txt"
S3, S3B, S4 = "3. Guardrail Research Table", "3b. NeMo Rail Inventory", "4. Candidate Comparison Groups"
S3C = "3c. NeMo Evaluation Tooling"
S3D = "3d. Llama Guard Inventory"
S3E = "3e. GovTech Sentinel Inventory"
ORDER = [S3, S3B, S3C, S3D, S3E, S4]
BAK4 = DIR + r"\AI Guardrails Research and Comparison.v4.xlsx"
BAK5 = DIR + r"\AI Guardrails Research and Comparison.v5.xlsx"
BAK6 = DIR + r"\AI Guardrails Research and Comparison.v6.xlsx"
BAK7 = DIR + r"\AI Guardrails Research and Comparison.v7.xlsx"
SKILL = (r"C:\Users\cys-c\AppData\Roaming\Claude\local-agent-mode-sessions\skills-plugin"
         r"\83a0dd2b-9f62-42a2-bdbb-7ca10dbfee57\8e055232-48fa-4a7d-80fb-77845d6c51b7\skills\xlsx")
R1, RN = 4, 83  # surface table rows (header is row 3)
thin = Side(style="thin", color="8EA9DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
AMBER = PatternFill("solid", fgColor="FFF2CC")
HFILL = PatternFill("solid", fgColor="4472C4")
WIDTHS = [5, 20, 11, 34, 11, 11, 45, 40, 40, 50]
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
XMLSP = "{http://www.w3.org/XML/1998/namespace}space"


def nemo_cols(ws3):
    """Sheet-3 column indexes of the NeMo product columns, found by the row-3 header prefix (F..U = 6..21)."""
    cols = [c for c in range(1, ws3.max_column + 1)
            if str(ws3.cell(row=3, column=c).value or "").startswith("NeMo Guardrails:")]
    assert len(cols) == 16 and cols == list(range(cols[0], cols[0] + 16)), cols
    return cols


def clean(t):
    return t.strip().replace("`", "").replace("**", "")


def parse_section(lines, start_pred):
    """Return (header, rows) of the first md table after the first line matching start_pred."""
    i = next(k for k, l in enumerate(lines) if start_pred(l))
    j = next(k for k in range(i + 1, len(lines)) if lines[k].startswith("|"))
    tbl = []
    while j < len(lines) and lines[j].startswith("|"):
        tbl.append([clean(c) for c in lines[j].strip().strip("|").split("|")])
        j += 1
    hdr, rows = tbl[0], tbl[2:]
    for r in rows:
        assert len(r) == len(hdr), (len(r), len(hdr), r[:3])
    return hdr, rows


def parse_inventory(path=INV):
    lines = open(path, encoding="utf-8").read().splitlines()
    surf = parse_section(lines, lambda l: l.startswith("## Output 1"))
    rtype = parse_section(lines, lambda l: l.startswith("## Output 2"))
    dual = None
    if any(l.startswith("## Dual-mapping evidence") for l in lines):
        dual = parse_section(lines, lambda l: l.startswith("## Dual-mapping evidence"))
    assert len(surf[1]) == 80 and len(rtype[1]) == 7, (len(surf[1]), len(rtype[1]))
    assert len(surf[0]) == 10
    return surf, rtype, dual


def _hdr(ws, r, hdr):
    for ci, h in enumerate(hdr, 1):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill, c.border = HFILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")


def _body(ws, r, row, center=(), amber_col=None):
    amber = amber_col is not None and row[amber_col].strip().lower() == "not covered"
    for ci, v in enumerate(row, 1):
        c = ws.cell(row=r, column=ci, value=v)
        c.font = Font(name="Arial", size=10)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top",
                                horizontal="center" if ci in center else None)
        if amber:
            c.fill = AMBER
    return r + 1


def add_sheet(wb, ws3):
    (sh, srows), (rh, rrows), dual = parse_inventory()
    if S3B in wb.sheetnames:
        wb.remove(wb[S3B])
    ws = wb.create_sheet(S3B, wb.sheetnames.index(S3) + 1)
    ws["A1"] = "3b. NeMo Guardrails Rail Inventory (v0.24.1)"
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    ws["A2"] = ("Source: NVIDIA Rail Engine Support + Engine Feature Support pages; GitHub NVIDIA-NeMo/Guardrails "
                "v0.24.1. Covered-by maps each surface to a Table 3 column.")
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="808080")
    _hdr(ws, 3, sh)
    r = 4
    for row in srows:
        r = _body(ws, r, row, center=(5, 6), amber_col=8)
    assert r == 84
    ws.auto_filter.ref = "A3:J83"
    r += 1  # blank row
    ws.cell(row=r, column=1, value="Rail types (user-defined / framework)").font = Font(name="Arial", size=10, bold=True)
    r += 1
    _hdr(ws, r, rh)
    r += 1
    rc = rh.index("Covered by Table 3 column")
    rt1 = r
    for row in rrows:
        r = _body(ws, r, row, center=(3, 4), amber_col=rc)
    rtN = r - 1
    if dual:
        r += 1
        ws.cell(row=r, column=1, value="Dual-mapping evidence (input surfaces only)").font = Font(name="Arial", size=10, bold=True)
        r += 1
        _hdr(ws, r, dual[0])
        r += 1
        for row in dual[1]:
            r = _body(ws, r, row)
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    add_panel(ws, ws3, r + 2, rt1, rtN, rc + 1)  # 2 blank rows after the last block
    return ws


# ---------------------------------------------------------------- formula summary panel (below the last block)
DIRS = ["input", "output", "retrieval", "tool_output", "tool_input"]
BUILTIN_ROWS = [  # (label, formula template with {H})
    ("built-in", '=COUNTIF({H},"built-in*")'),
    ("third-party", '=COUNTIF({H},"third-party*")'),
    ("of which commercial vendor", '=COUNTIF({H},"*commercial vendor*")'),
    ("of which licensing to be verified", '=COUNTIF({H},"*licensing to be verified*")'),
    ("of which open-source / open-weight",
     '=SUMPRODUCT(--((ISNUMBER(SEARCH("open-source",{H}))+ISNUMBER(SEARCH("open-weight",{H})))>0))'),
    ("deprecated", '=COUNTIF({H},"*deprecated*")'),
]
PANEL = {}  # filled by add_panel: row anchors and rail-type table location
LBL_LAST = 4  # labels occupy A:D (merged); numbers start in column E
NUM0 = 5


def _rng(col):
    return f"${col}${R1}:${col}${RN}"


def _pcell(ws, r, c, value, hdr=False, num=False, bold=False):
    x = ws.cell(row=r, column=c, value=value)
    x.border = BORDER
    if hdr:
        x.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        x.fill = HFILL
        x.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c >= NUM0 else None)
    else:
        x.font = Font(name="Arial", size=10, bold=bold)
        x.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if num else None)
    return x


def _plabel(ws, r, text, hdr=False, bold=False):
    """Label in A, merged A:D."""
    for c in range(1, LBL_LAST + 1):
        _pcell(ws, r, c, text if c == 1 else None, hdr=hdr, bold=bold)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=LBL_LAST)


def _ptitle(ws, r, text, last_col):
    for c in range(1, last_col + 1):
        _pcell(ws, r, c, text if c == 1 else None, hdr=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=last_col)


def add_panel(ws, ws3, start, rt1, rtN, rcol):
    """start = row of the bold panel heading; rail-type table rows rt1..rtN, Covered-by column index rcol."""
    C, E, F, H, I = (_rng(x) for x in "CEFHI")
    RI = f"${get_column_letter(rcol)}${rt1}:${get_column_letter(rcol)}${rtN}"
    assert ws.cell(row=R1 - 1, column=3).value == "Direction" and ws.cell(row=RN + 1, column=1).value is None \
        and ws.cell(row=RN, column=1).value is not None
    ws.cell(row=start, column=1, value="Summary panel (live formulas)").font = Font(name="Arial", size=11, bold=True)
    P = PANEL
    P.update(head=start, A_title=start + 1, A_hdr=start + 2, A_first=start + 3, A_total=start + 8,
             B_title=start + 10, B_hdr=start + 11, B_first=start + 12,
             C_title=start + 19, C_hdr=start + 20, C_first=start + 21, C_nc=start + 37,
             rt=(rt1, rtN, rcol), RI=RI)
    n = NUM0
    # Panel A
    _ptitle(ws, P["A_title"], "Surfaces by direction × engine", n + 3)
    _plabel(ws, P["A_hdr"], "Direction", hdr=True)
    for i, h in enumerate(["Surfaces", "LLMRails ✓", "IORails ✓", "IORails ✗"]):
        _pcell(ws, P["A_hdr"], n + i, h, hdr=True)
    for k, d in enumerate(DIRS):
        r = P["A_first"] + k
        _plabel(ws, r, d)
        _pcell(ws, r, n, f"=COUNTIFS({C},$A{r})", num=True)
        _pcell(ws, r, n + 1, f'=COUNTIFS({C},$A{r},{E},"✓*")', num=True)
        _pcell(ws, r, n + 2, f'=COUNTIFS({C},$A{r},{F},"✓*")', num=True)
        _pcell(ws, r, n + 3, f'=COUNTIFS({C},$A{r},{F},"✗*")', num=True)
    r = P["A_total"]
    _plabel(ws, r, "Total", bold=True)
    for c in range(n, n + 4):
        cl = get_column_letter(c)
        _pcell(ws, r, c, f"=SUM({cl}{P['A_first']}:{cl}{r - 1})", num=True, bold=True)
    # Panel B
    _ptitle(ws, P["B_title"], "Built-in vs third-party", n)
    _plabel(ws, P["B_hdr"], "Category", hdr=True)
    _pcell(ws, P["B_hdr"], n, "Surfaces", hdr=True)
    for k, (lab, f) in enumerate(BUILTIN_ROWS):
        r = P["B_first"] + k
        _plabel(ws, r, lab)
        _pcell(ws, r, n, f.format(H=H), num=True)
    # Panel C
    _ptitle(ws, P["C_title"], "Surfaces and rail types per Table 3 column", n + 2)
    _plabel(ws, P["C_hdr"], "Table 3 column (NeMo)", hdr=True)
    for i, h in enumerate(["Surfaces", "Rail-type rows", "Total"]):
        _pcell(ws, P["C_hdr"], n + i, h, hdr=True)
    for k, col in enumerate(nemo_cols(ws3)):
        r = P["C_first"] + k
        hv = ws3.cell(row=3, column=col).value
        assert "*" not in hv and "?" not in hv and "~" not in hv and len(hv) < 250, hv  # no wildcard escapes needed
        _plabel(ws, r, f"='{S3}'!{get_column_letter(col)}3")
        _pcell(ws, r, n, f'=COUNTIF({I},"*"&$A{r}&"*")', num=True)
        _pcell(ws, r, n + 1, f'=COUNTIF({RI},"*"&$A{r}&"*")', num=True)
        _pcell(ws, r, n + 2, f"={get_column_letter(n)}{r}+{get_column_letter(n + 1)}{r}", num=True)
    assert P["C_first"] + 16 == P["C_nc"]
    r = P["C_nc"]
    _plabel(ws, r, "not covered", bold=True)
    _pcell(ws, r, n, f'=COUNTIF({I},"not covered")', num=True, bold=True)
    _pcell(ws, r, n + 1, f'=COUNTIF({RI},"not covered")', num=True, bold=True)
    _pcell(ws, r, n + 2, f"={get_column_letter(n)}{r}+{get_column_letter(n + 1)}{r}", num=True, bold=True)
    return P["C_nc"]


def python_panel(rows, hdrs, rail_cov):
    """Expected panel values computed in Python from the surface rows (list of lists A..J) and Table 3 headers."""
    a = {}
    for d in DIRS:
        sel = [r for r in rows if r[2] == d]
        a[d] = (len(sel), sum(str(r[4]).startswith("✓") for r in sel), sum(str(r[5]).startswith("✓") for r in sel),
                sum(str(r[5]).startswith("✗") for r in sel))
    a["Total"] = tuple(sum(a[d][i] for d in DIRS) for i in range(4))
    h = [str(r[7]).lower() for r in rows]
    b = [sum(x.startswith("built-in") for x in h), sum(x.startswith("third-party") for x in h),
         sum("commercial vendor" in x for x in h), sum("licensing to be verified" in x for x in h),
         sum(("open-source" in x) or ("open-weight" in x) for x in h), sum("deprecated" in x for x in h)]
    c = [sum(hd in [v.strip() for v in str(r[8]).split("; ")] for r in rows) for hd in hdrs]
    nc = sum(str(r[8]).strip().lower() == "not covered" for r in rows)
    # rail-type Covered-by cells hold annotated text (e.g. "... (new column)", "also X (omit)"), so the
    # wildcard "*header*" in the formula is a case-insensitive substring match; mirror that here.
    cr = [sum(hd.lower() in str(x).lower() for x in rail_cov) for hd in hdrs]
    ncr = sum(str(x).strip().lower() == "not covered" for x in rail_cov)
    return a, b, c, nc, cr, ncr


def hygiene(xlsx):
    """XML hygiene on a saved xlsx. Returns (unpreserved_ws_t, empty_t, r_without_rPr, nparts_checked, well_formed)."""
    from lxml import etree
    z = zipfile.ZipFile(xlsx)
    ssp = norpr = emp = 0
    wf = True
    for n in z.namelist():
        if n.endswith((".xml", ".rels")):
            try:
                etree.fromstring(z.read(n))
            except Exception:
                wf = False
    parts = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet") or n == "xl/sharedStrings.xml"]
    for n in parts:
        root = etree.fromstring(z.read(n))
        for t in root.iter(NS + "t"):
            tx = t.text or ""
            emp += tx == ""
            ssp += (tx != tx.strip() and t.get(XMLSP) != "preserve")
        norpr += sum(1 for x in root.iter(NS + "r") if x.find(NS + "rPr") is None)
    return ssp, emp, norpr, len(parts), wf


def verify_panel(xlsx, post_check=None):
    from openpyxl import load_workbook
    print("\n=== 3b formula panel verification ===")
    P = PANEL
    wf = load_workbook(xlsx)
    ws = wf[S3B]
    hdrs = [wf[S3].cell(row=3, column=c).value for c in nemo_cols(wf[S3])]
    rows = [[c.value for c in ws[r][:10]] for r in range(R1, RN + 1)]
    print("surface table: header row 3 =", ws.cell(row=3, column=1).value, "..; last row", RN, "=", ws.cell(row=RN, column=1).value,
          "; row", RN + 1, "=", ws.cell(row=RN + 1, column=1).value, "(None); rows:", len(rows), "(80)")
    rt1, rtN, rcol = P["rt"]
    print("rail-type table rows", rt1, "-", rtN, "Covered-by column", get_column_letter(rcol), "; panel heading row",
          P["head"], "col A =", repr(ws.cell(row=P["head"], column=1).value))
    pcells = [c for row in ws.iter_rows(min_row=P["head"]) for c in row if c.value is not None]
    inside = [c.coordinate for c in pcells if c.row <= RN]
    print("panel cells:", len(pcells), "; rows", P["head"], "-", P["C_nc"], "; panel cells inside A3:J83:", len(inside))
    assert P["head"] > RN and not inside
    gap = [r for r in range(P["head"] - 2, P["head"]) if any(c.value is not None for c in ws[r])]
    print("two rows above the heading empty:", not gap)
    nf = sum(1 for c in pcells if isinstance(c.value, str) and c.value.startswith("="))
    banned = [c.coordinate for c in pcells
              if isinstance(c.value, str) and re.search(r"FILTER|UNIQUE|XLOOKUP|LET\(|SORT|SEQUENCE", c.value)]
    print("panel formulas:", nf, " dynamic-array functions found:", banned)
    rail_cov = [ws.cell(row=r, column=rcol).value for r in range(rt1, rtN + 1)]
    print("rail-type Covered-by values:", rail_cov)
    ea, eb, ec, enc, ecr, encr = python_panel(rows, hdrs, rail_cov)
    print("Python-expected  A:", ea)
    print("Python-expected  B:", dict(zip([x for x, _ in BUILTIN_ROWS], eb)))
    print("Python-expected  C surfaces:", dict(zip(hdrs, ec)), " not covered:", enc)
    print("Python-expected  C rail-type rows:", dict(zip(hdrs, ecr)), " not covered:", encr)
    print("Python-expected  C total:", dict(zip(hdrs, [a + b for a, b in zip(ec, ecr)])), " not covered:", enc + encr)

    if shutil.which("soffice") is None:
        print("LibreOffice (soffice) is NOT available on this machine: recalc skipped; formulas have no cached "
              "values (Excel will calculate them on open). Panel values above are Python-expected, NOT recalculated.")
        return None
    tmp = xlsx + ".pre_recalc.tmp"
    shutil.copy2(xlsx, tmp)
    env = dict(os.environ)
    r = subprocess.run([sys.executable, os.path.join(SKILL, "scripts", "recalc.py"), xlsx, "60"],
                       capture_output=True, text=True, env=env, cwd=os.path.join(SKILL, "scripts"))
    print("recalc stdout:", r.stdout.strip(), "\nrecalc stderr:", r.stderr.strip()[-500:])
    try:
        res = json.loads(r.stdout)
    except Exception:
        res = {}
    ok = res.get("status") == "success" and res.get("total_errors") == 0
    print("recalc status:", res.get("status"), " errors:", res.get("total_errors"), "->", "OK" if ok else "FAIL")
    if not ok:
        shutil.copy2(tmp, xlsx)
        os.remove(tmp)
        print("recalc failed; restored the openpyxl-written file (formulas uncached).")
        return False
    wv = load_workbook(xlsx, data_only=True)[S3B]
    g = lambda r, c: wv.cell(row=r, column=c).value
    A = {d: tuple(g(P["A_first"] + k, NUM0 + j) for j in range(4)) for k, d in enumerate(DIRS)}
    A["Total"] = tuple(g(P["A_total"], NUM0 + j) for j in range(4))
    B = [g(P["B_first"] + k, NUM0) for k in range(6)]
    Cv = [(g(P["C_first"] + k, 1), g(P["C_first"] + k, NUM0), g(P["C_first"] + k, NUM0 + 1),
           g(P["C_first"] + k, NUM0 + 2)) for k in range(16)]
    NC = g(P["C_nc"], NUM0)
    NCR, NCT = g(P["C_nc"], NUM0 + 1), g(P["C_nc"], NUM0 + 2)
    print("recalculated A:", A, "\nrecalculated B:", dict(zip([x for x, _ in BUILTIN_ROWS], B)),
          "\nrecalculated C:", Cv, " not covered:", NC)
    chk = {
        "A total surfaces 80": A["Total"][0] == 80,
        "input/output/retrieval 32/35/11": (A["input"][0], A["output"][0], A["retrieval"][0]) == (32, 35, 11),
        "IORails ✓ in/out/retr 59": sum(A[d][2] for d in DIRS[:3]) == 59,
        "A equals Python": A == ea,
        "C labels equal sheet-3 headers": [x[0] for x in Cv] == hdrs,
        "C surface counts equal Python": [x[1] for x in Cv] == ec,
        "C rail-type counts equal Python": [x[2] for x in Cv] == ecr,
        "C totals equal sum": [x[3] for x in Cv] == [a + b for a, b in zip(ec, ecr)],
        "not covered 0 (surfaces, rail types, total)": (NC, NCR, NCT) == (0, 0, 0) == (enc, encr, enc + encr),
        "B built-in + third-party = 80": B[0] + B[1] == 80,
        "B equals Python": B == eb,
    }
    for k, v in chk.items():
        print("  check:", k, "OK" if v else "FAIL")
    # LibreOffice rewrote the file: re-run hygiene + sheet-3 / rich-text checks
    ssp, emp, norpr, n, wfm = hygiene(xlsx)
    print("post-recalc XML hygiene: unpreserved-ws <t>=%d, empty <t>=%d, <r> without rPr=%d, parts checked=%d, "
          "all parts well-formed=%s" % (ssp, emp, norpr, n, wfm))
    intact = ssp == 0 and emp == 0 and norpr == 0 and wfm
    if post_check is not None:
        intact = post_check(xlsx) and intact
    if not intact:
        shutil.copy2(tmp, xlsx)
        print("LibreOffice rewrite DAMAGED rich text / styles: NOT shipping the recalculated file; restored the "
              "openpyxl-written file (formulas have no cached values; Excel calculates them on open).")
    else:
        print("recalculated file intact (rich text and styles survived); shipping it.")
    os.remove(tmp)
    return intact


URL_RE = re.compile(r"https?://[^\s)\]>\"'<,;]+")


def verify(xlsx, wb2, bak, post_check=None):
    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from lxml import etree
    print("\n=== build_inventory verification ===")
    exp = ORDER
    print("sheet order:", wb2.sheetnames, "OK" if wb2.sheetnames == exp else "FAIL")
    ws3, ws = wb2[S3], wb2[S3B]
    rows = [[c.value for c in ws[r][:10]] for r in range(4, 84)]
    print("surface rows:", len(rows), "(80); header:", [c.value for c in ws[3][:10]])
    dc = Counter(r[2] for r in rows)
    want = {"input": 32, "output": 35, "retrieval": 11, "tool_output": 1, "tool_input": 1}
    print("direction counts:", dict(dc), "OK" if dict(dc) == want else "FAIL (want %s)" % want)
    io = sum(1 for r in rows if r[2] in ("input", "output", "retrieval") and str(r[5]).startswith("✓"))
    print("IORails ✓ among input/output/retrieval:", io, "(59)", "OK" if io == 59 else "FAIL")
    hdr3 = {ws3.cell(row=3, column=c).value for c in nemo_cols(ws3)}
    bad = []
    for i, r in enumerate(rows):
        for v in str(r[8]).split(";"):
            v = v.strip()
            if v != "not covered" and v not in hdr3:
                bad.append((i + 4, v))
    print("covered-by values not exact sheet-3 header:", len(bad), bad[:5])
    rt = [r for r in range(85, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Rail type"]
    names = []
    if rt:
        k = rt[0] + 1
        while ws.cell(row=k, column=1).value:
            names.append(ws.cell(row=k, column=1).value)
            k += 1
    print("rail-type header row:", rt, "; rail-type rows:", len(names), "(7)", names)
    mk = [c.coordinate for row in ws.iter_rows() for c in row
          if isinstance(c.value, str) and ("**" in c.value or "`" in c.value)]
    print("cells with ** or backtick on 3b:", len(mk), mk[:5])
    print("freeze:", ws.freeze_panes, " autofilter:", ws.auto_filter.ref,
          " widths:", [ws.column_dimensions[get_column_letter(i)].width for i in range(1, 11)])
    print("max_row 3b:", ws.max_row)

    # sheet 4 is verified by build_groups_sheet.verify (it holds the groups table and panel since v7)

    z = zipfile.ZipFile(xlsx)
    ssp = norpr = emp = 0
    parts = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet") or n == "xl/sharedStrings.xml"]
    for n in z.namelist():
        if n.endswith((".xml", ".rels")):
            etree.fromstring(z.read(n))
    for n in parts:
        root = etree.fromstring(z.read(n))
        for t in root.iter(NS + "t"):
            tx = t.text or ""
            emp += tx == ""
            ssp += (tx != tx.strip() and t.get(XMLSP) != "preserve")
        norpr += sum(1 for x in root.iter(NS + "r") if x.find(NS + "rPr") is None)
    print("XML (%d parts checked, all parts well-formed): t-ws-no-preserve=%d, r-no-rPr=%d, empty-t=%d"
          % (len(parts), ssp, norpr, emp))

    urls, seen = [], set()

    def add(text):
        for u in URL_RE.findall(text or ""):
            u = u.rstrip(".:")
            if u not in seen:
                seen.add(u)
                urls.append(u)
    for c in nemo_cols(ws3):
        v = ws3.cell(row=21, column=c).value
        add("".join(b.text if isinstance(b, TextBlock) else str(b) for b in v)
            if isinstance(v, CellRichText) else v)
    n9 = len(urls)
    for r in rows:
        add(r[9])
    out = "\n".join(urls) + "\n"
    for attempt in range(3):
        try:
            with open(URLS_OUT, "w", encoding="utf-8", newline="\n") as f:
                f.write(out)
            break
        except PermissionError:
            print("PermissionError writing urls; retry in 15s")
            time.sleep(15)
    else:
        print("FAILED to write", URLS_OUT)
    print("urls: %d unique (%d from sheet3 R9 detail, %d new from 3b) -> %s"
          % (len(urls), n9, len(urls) - n9, URLS_OUT))
    verify_panel(xlsx, post_check)
