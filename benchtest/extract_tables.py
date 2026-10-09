"""Extract the 2 tables from the AI Guardrails docx into an xlsx (keeps bold runs).

Refuses to overwrite an existing workbook whose sheet 4 has been filled (A4 differs from the docx template text)
unless --force is passed: python extract_tables.py --force
"""
import os, sys
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DIR = r"C:\Users\cys-c\Desktop\gzqr\benchtest"
SRC = DIR + r"\AI Guardrails Research and Comparison.docx"
DST = DIR + r"\AI Guardrails Research and Comparison.xlsx"

FIX = lambda s: s.replace("\ufffd", "\u2019")

SHEETS = [
    # sheet name, full title, widths(fn col index->width), freeze
    ("3. Guardrail Research Table", "3. Guardrail Research Table",
     lambda i: {0: 6, 1: 28, 2: 32}.get(i, 48), "D4"),
    ("4. Candidate Comparison Groups", "4. Forming Candidate Comparison Groups",
     lambda i: 28 if i == 0 else 34, "B4"),
]

thin = Side(style="thin", color="8EA9DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
ALIGN = Alignment(wrap_text=True, vertical="top")


def cell_runs(cell):
    """Return list of (text, bold) with adjacent same-bold runs merged;
    paragraphs joined with '\\n'."""
    runs = []
    for pi, p in enumerate(cell.paragraphs):
        if pi:
            runs.append(("\n", None))  # None = inherits neighbour's state
        for r in p.runs:
            runs.append((FIX(r.text), bool(r.bold)))
    # resolve separators: attach to previous run's bold state
    resolved = []
    for t, b in runs:
        if b is None:
            b = resolved[-1][1] if resolved else False
        resolved.append((t, b))
    merged = []
    for t, b in resolved:
        if merged and merged[-1][1] == b:
            merged[-1][0] += t
        else:
            merged.append([t, b])
    return merged


def plain_text(cell):
    return "\n".join(FIX(p.text) for p in cell.paragraphs)


def cell_value(cell):
    merged = cell_runs(cell)
    if not any(b for _, b in merged):
        return plain_text(cell)
    return CellRichText(*[
        TextBlock(InlineFont(b=True, rFont="Arial", sz=10) if b
                  else InlineFont(rFont="Arial", sz=10), t)
        for t, b in merged])


def guard_sheet4(doc, force=False):
    """SystemExit if DST exists and its sheet 4 A4 is not the docx template text (sheet 4 was filled), unless force."""
    if force or not os.path.exists(DST):
        return
    name4 = SHEETS[1][0]
    wb = load_workbook(DST)
    if name4 not in wb.sheetnames:
        return
    cur = wb[name4]["A4"].value
    tmpl = plain_text(doc.tables[1].rows[1].cells[0])
    if cur != tmpl:
        raise SystemExit("Refusing to overwrite %s: sheet 4 A4 is %r, not the docx template text %r "
                         "(sheet 4 has been filled). Pass --force to overwrite anyway." % (DST, str(cur)[:60], tmpl))


def build(force=False):
    doc = Document(SRC)
    guard_sheet4(doc, force)
    wb = Workbook()
    wb.remove(wb.active)
    for tbl, (name, title, width, freeze) in zip(doc.tables, SHEETS):
        ws = wb.create_sheet(name)
        ws["A1"] = title
        ws["A1"].font = Font(name="Arial", size=13, bold=True)
        ncols = len(tbl.columns)
        for ri, row in enumerate(tbl.rows):
            r = ri + 3
            for ci, c in enumerate(row.cells):
                xc = ws.cell(row=r, column=ci + 1)
                if ri == 0:
                    xc.value = plain_text(c)
                    xc.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                    xc.fill = PatternFill("solid", fgColor="4472C4")
                else:
                    xc.value = cell_value(c)
                    xc.font = Font(name="Arial", size=10)
                    if ri % 2 == 1:
                        xc.fill = PatternFill("solid", fgColor="DCE6F2")
                xc.alignment = ALIGN
                xc.border = BORDER
        for ci in range(ncols):
            ws.column_dimensions[get_column_letter(ci + 1)].width = width(ci)
        ws.freeze_panes = freeze
    wb.save(DST)
    return doc


def verify(doc):
    wb = load_workbook(DST, rich_text=True)
    mism = 0
    total = 0
    for tbl, (name, *_rest) in zip(doc.tables, SHEETS):
        ws = wb[name]
        for ri, row in enumerate(tbl.rows):
            for ci, c in enumerate(row.cells):
                total += 1
                got = ws.cell(row=ri + 3, column=ci + 1).value
                got = "" if got is None else str(got)
                if got != plain_text(c):
                    mism += 1
                    print("MISMATCH", name, ri + 3, ci + 1, repr(got), repr(plain_text(c)))
    bad = [(ws.title, c.coordinate) for ws in wb for row in ws.iter_rows(min_row=3)
           for c in row if c.font.name != "Arial" or c.font.sz != 10]
    print("non-Arial/10 cells from row 3:", len(bad), bad[:5])
    print(f"cells compared: {total}, mismatch count: {mism}")
    ws = wb["3. Guardrail Research Table"]
    for ref in ("D4", "E4"):
        v = ws[ref].value
        if isinstance(v, CellRichText):
            print(ref, [(b.text, bool(b.font.b)) if isinstance(b, TextBlock) else (b, False)
                        for b in v])
        else:
            print(ref, "plain:", repr(v))


if __name__ == "__main__":
    verify(build(force="--force" in sys.argv[1:]))
