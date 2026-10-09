"""Independent verification of the Presidio apply (sheet 3 AH-AM, sheet 3f, all other sheets vs the previous commit).

Baseline = the workbook at git HEAD (the previous commit), not a frozen file under baselines/.
"""
import os
import subprocess
import sys
import tempfile
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils import get_column_letter

import build_two_level as B2
import build_inventory as BI
import build_presidio_inventory as BP
import inventory_sheet as IS
from fill_nemo_columns import LABEL_RE

XLSX = B2.XLSX
URLS = str(BI.P.DRAFTS / "presidio_urls.txt")
S3, S3B, S3C, S3D, S3E, S3F, S4 = B2.S3, BI.S3B, BI.S3C, BI.S3D, BI.S3E, BP.S3F, B2.S4
RESULTS = []


def check(name, ok, extra=""):
    RESULTS.append(bool(ok))
    print(("PASS" if ok else "FAIL"), "-", name, extra)


plain = B2.plain_of


def runs_of(v):
    return [(b.text, bool(b.font.b)) for b in v] if isinstance(v, CellRichText) else [(plain(v), False)]


def _col(c):
    return None if c is None else (c.rgb if c.type == "rgb" else (c.type, c.theme if c.type == "theme" else c.indexed))


def style_of(c):
    b = c.border
    return (c.font.name, c.font.sz, c.font.b, c.font.i, _col(c.font.color), c.fill.fill_type, _col(c.fill.fgColor),
            c.alignment.wrap_text, c.alignment.vertical, c.alignment.horizontal,
            tuple((s.style, _col(s.color)) if s is not None else None for s in (b.left, b.right, b.top, b.bottom)))


def full_snap(w):
    cells = [(c.coordinate, plain(c.value), runs_of(c.value), style_of(c)) for row in w.iter_rows() for c in row]
    layout = (sorted(str(m) for m in w.merged_cells.ranges), w.freeze_panes, w.auto_filter.ref,
              {k: v.width for k, v in w.column_dimensions.items()},
              {k: (v.height, v.outline_level, v.hidden) for k, v in w.row_dimensions.items()},
              (w.max_row, w.max_column))
    return cells, layout


# ---------------------------------------------------------------- previous commit's workbook
prev = os.path.join(tempfile.mkdtemp(), "prev.xlsx")
with open(prev, "wb") as f:
    f.write(subprocess.run(["git", "show", 'HEAD:benchtest/AI Guardrails Research and Comparison.xlsx'],
                           capture_output=True, check=True, cwd=str(BI.P.ROOT)).stdout)
new = load_workbook(XLSX, rich_text=True)
old = load_workbook(prev, rich_text=True)
print("previous-commit sheet names:", old.sheetnames)
check("sheet order", new.sheetnames == list(BI.ORDER), str(new.sheetnames))
check("previous sheets all present, in order, plus only 3f",
      [n for n in new.sheetnames if n != S3F] == old.sheetnames == [S3, S3B, S3C, S3D, S3E, S4])

# ---------------------------------------------------------------- ranges from drafts
mds = [B2.parse_md(m) for m in B2.MDS]
n_each = [len(m) for m in mds]
assert len(mds) == 4, n_each
OLD_LAST = 5 + sum(n_each[:3])            # AG = 33
PD0 = OLD_LAST + 1                        # AH = 34
md_pd = mds[3]
LAST = PD0 + len(md_pd) - 1               # AM = 39
print("derived: old last col %s (%d), Presidio %s-%s (%d), last col %d" % (
    get_column_letter(OLD_LAST), OLD_LAST, get_column_letter(PD0), get_column_letter(LAST), len(md_pd), LAST))

ws, wo = new[S3], old[S3]
check("sheet 3: %d columns (A-%s)" % (LAST, get_column_letter(LAST)), ws.max_column == LAST == 39, str(ws.max_column))
check("sheet 3: max_row 21", ws.max_row == 21, str(ws.max_row))
check("27 merges equal to previous commit",
      len(ws.merged_cells.ranges) == 27 and sorted(map(str, ws.merged_cells.ranges)) == sorted(map(str, wo.merged_cells.ranges)))
lv = [r for r in range(1, 23) if ws.row_dimensions[r].outline_level == 1]
check("outline levels (rows 5,7,..,21 level 1)", lv == list(range(5, 22, 2)), str(lv))
check("freeze panes E4 (as previous)", ws.freeze_panes == wo.freeze_panes == "E4")
check("summaryBelow False", ws.sheet_properties.outlinePr.summaryBelow is False)
widths = [ws.column_dimensions[get_column_letter(i)].width for i in range(1, LAST + 1)]
check("column widths 6/28/32/10 then 48", widths == [6, 28, 32, 10] + [48] * (LAST - 4), str(widths[-3:]))
check("old column widths A-%s identical to previous commit" % get_column_letter(OLD_LAST),
      all(ws.column_dimensions[get_column_letter(i)].width == wo.column_dimensions[get_column_letter(i)].width
          for i in range(1, OLD_LAST + 1)))
pd_hdrs = [c["header"] for c in md_pd]
got_h = [ws.cell(row=3, column=c).value for c in range(PD0, LAST + 1)]
check("AH3..AM3 are the 6 Presidio headers in md order (prefix 'Presidio:')",
      got_h == pd_hdrs and len(pd_hdrs) == 6 and all(h.startswith("Presidio:") for h in pd_hdrs),
      str([h[10:45] for h in pd_hdrs]))
check("old header row identical (A3:AG3)", [ws.cell(row=3, column=c).value for c in range(1, OLD_LAST + 1)]
      == [wo.cell(row=3, column=c).value for c in range(1, OLD_LAST + 1)])

# ---------------------------------------------------------------- A-AG vs previous commit (all rows)
diff = sum(plain(ws.cell(row=r, column=c).value) != plain(wo.cell(row=r, column=c).value)
           for r in range(1, 23) for c in range(1, OLD_LAST + 1))
check("A-AG plain text vs previous commit: 0 diffs", diff == 0, f"({diff} diffs)")
rdiff = sum(runs_of(ws.cell(row=r, column=c).value) != runs_of(wo.cell(row=r, column=c).value)
            for r in range(4, 22) for c in range(6, OLD_LAST + 1))
check("F-AG rows 4-21 rich-text runs (text+bold) vs previous commit: 0 diffs", rdiff == 0, f"({rdiff} diffs)")
sdiff = sum(style_of(ws.cell(row=r, column=c)) != style_of(wo.cell(row=r, column=c))
            for r in range(1, 23) for c in range(1, OLD_LAST + 1))
check("A-AG cell styles vs previous commit: 0 diffs", sdiff == 0, f"({sdiff} diffs)")

# ---------------------------------------------------------------- Presidio AH-AM vs md
mism = ncells = 0
for n in range(1, 10):
    sr, dr = 4 + 2 * (n - 1), 5 + 2 * (n - 1)
    for ci, col in enumerate(md_pd):
        cc = PD0 + ci
        d = col["R"][n]
        _, sp = B2.rich(d["summary"], 10, "v")
        _, dp = B2.rich("\n".join(d["detail"]), 9, "v")
        mism += plain(ws.cell(row=sr, column=cc).value) != sp
        mism += plain(ws.cell(row=dr, column=cc).value) != dp
        ncells += 2
check("AH-AM plain text vs presidio_two_level.md parse: 0 mismatches (%d cells)" % ncells,
      mism == 0 and ncells == 108, f"({mism})")
fails = nsum = 0
for n in range(1, 10):
    sr = 4 + 2 * (n - 1)
    for cc in range(PD0, LAST + 1):
        nsum += 1
        c = ws.cell(row=sr, column=cc)
        blocks = list(c.value) if isinstance(c.value, CellRichText) else [c.value]
        p = plain(c.value)
        probs = []
        if n <= 8 and not (isinstance(blocks[0], TextBlock) and blocks[0].font.b):
            probs.append("first-not-bold")
        if n <= 7 and not any(isinstance(b, TextBlock) and b.font.b and LABEL_RE.search(b.text) for b in blocks):
            probs.append("no-bold-label")
        if any(s in p for s in ("**", "`", "_", "$")):
            probs.append("marker")
        if len(p.split()) > 60:
            probs.append("words=%d" % len(p.split()))
        if "\n" in p:
            probs.append("newline")
        if probs:
            fails += 1
            print("   ", c.coordinate, probs)
check("AH-AM Summary style checks (%d cells)" % nsum, fails == 0 and nsum == 54, f"({fails} failing)")
bad_md = [ws.cell(row=r, column=c).coordinate for r in range(4, 22) for c in range(PD0, LAST + 1)
          if any(s in plain(ws.cell(row=r, column=c).value) for s in ("**", "`"))]
check("AH-AM no ** or backticks (rows 4-21)", not bad_md, str(bad_md[:5]))
nonar = [ws.cell(row=r, column=c).coordinate for r in range(3, 22) for c in range(PD0, LAST + 1)
         if ws.cell(row=r, column=c).font.name != "Arial"]
rich_nonar = sum(1 for r in range(4, 22) for c in range(PD0, LAST + 1)
                 for b in (ws.cell(row=r, column=c).value if isinstance(ws.cell(row=r, column=c).value, CellRichText) else [])
                 if isinstance(b, TextBlock) and b.font.rFont != "Arial")
check("AH-AM non-Arial cells / rich runs: 0", not nonar and rich_nonar == 0, f"({nonar[:5]}, rich runs {rich_nonar})")
check("AH-AM cell styles equal the AG styles row by row, all 6 columns (same house format)",
      all(style_of(ws.cell(row=r, column=PD0 + k)) == style_of(ws.cell(row=r, column=OLD_LAST))
          for r in range(1, 22) for k in range(6)))

# ---------------------------------------------------------------- 3b, 3c, 3d, 3e, 4 identical to previous commit
for nm in (S3B, S3C, S3D, S3E, S4):
    a, al = full_snap(new[nm])
    b, bl = full_snap(old[nm])
    nd = sum(1 for x, y in zip(a, b) if x != y) + abs(len(a) - len(b))
    check(f"{nm}: values (incl. formulas), rich text, styles identical to previous commit", a == b,
          f"({len(a)} cells, {nd} diffs)")
    check(f"{nm}: merges, freeze, autofilter, widths, row heights/outline, extents identical", al == bl,
          "" if al == bl else str([(i, x, y) for i, (x, y) in enumerate(zip(al, bl)) if x != y])[:300])

# ---------------------------------------------------------------- 3f
w3f = new[S3F]
blocks = IS.parse_md(BP.CFG)
pos = [IS.block_rows(w3f, b["hdr"][0]) for b in blocks]
check("3f block row counts 14/31/10/14", [len(p[1]) for p in pos] == [14, 31, 10, 14], str([len(p[1]) for p in pos]))
mm = tot = 0
for b, (h, rows) in zip(blocks, pos):
    got = [[w3f.cell(row=r, column=c).value for c in range(1, len(b["hdr"]) + 1)] for r in [h] + rows]
    want = [b["hdr"]] + b["rows"]
    tot += sum(len(x) for x in want)
    mm += sum(1 for g, w in zip(got, want) for x, y in zip(g, w) if x != y) + abs(len(got) - len(want))
check("3f every cell equals cleaned md text (%d cells)" % tot, mm == 0, f"({mm} mismatches)")
h0, r0 = pos[0]
check("3f autofilter = block (a) only", w3f.auto_filter.ref == f"A{h0}:H{r0[-1]}", w3f.auto_filter.ref)
hdr3 = {ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)}
badc, ncov, amber_bad = [], 0, 0
for b, (h, rows) in zip(blocks, pos):
    cc = b["hdr"].index(BP.COV) + 1
    for r in rows:
        v = w3f.cell(row=r, column=cc)
        mk = v.value in (BP.LEGACY, BP.INV_ONLY)
        amber_bad += mk != (v.fill.fill_type == "solid" and str(v.fill.fgColor.rgb).endswith("FFF2CC"))
        for part in str(v.value).split(";"):
            part = part.strip()
            ncov += 1
            if part not in hdr3 and part not in (BP.LEGACY, BP.INV_ONLY):
                badc.append((r, part))
check("3f Covered-by values valid (real sheet-3 header or legacy/inventory-only marker), all 4 blocks", not badc,
      f"({ncov} values; bad {badc[:3]})")
check("3f amber fill exactly on marker cells", amber_bad == 0)
ctx = dict(ws=w3f, ws3=ws, cfg=BP.CFG, blocks=blocks, pos=pos, hdr3=hdr3)
for fn in BP.CFG["validators"]:
    ok_, msgs = fn(ctx)
    check("3f validator " + fn.__name__, ok_, "| " + " ".join(msgs)[:300])
ok_p, counts, totals = IS.verify_panel(new, BP.CFG, w3f, ws, blocks, pos)
check("3f coverage panel equals spec; every Presidio column covered; inventory-only 4, legacy 1",
      ok_p and totals["Rows marked inventory only"] == 4 and totals["Rows marked legacy"] == 1, str(totals))
mk = [c.coordinate for row in w3f.iter_rows() for c in row if isinstance(c.value, str) and ("**" in c.value or "`" in c.value)]
check("3f no ** or backticks", not mk, str(mk[:5]))
nonar = [c.coordinate for row in w3f.iter_rows() for c in row
         if type(c).__name__ != "MergedCell" and c.value is not None and c.font.name != "Arial"]
from build_eval_sheet import xml_non_arial
nx, bx = xml_non_arial(XLSX, S3F)
check("3f all Arial (cells + raw XML <c> check)", not nonar and bx == 0, f"(openpyxl {len(nonar)}, raw XML {bx} of {nx})")
check("3f A1 bold 13; A2 italic grey; freeze A4",
      w3f["A1"].font.b and w3f["A1"].font.sz == 13 and w3f["A2"].font.i and w3f["A2"].font.color.rgb.endswith("808080")
      and w3f.freeze_panes == "A4")

# ---------------------------------------------------------------- XML hygiene
ssp, emp, norpr, n, wf = BI.hygiene(XLSX)
check("XML hygiene (all parts well-formed; worksheets + sharedStrings)", ssp == 0 and emp == 0 and norpr == 0 and wf,
      f"(unpreserved <t>={ssp}, <r> without rPr={norpr}, empty <t>={emp}, parts checked={n}, well-formed={wf})")

# ---------------------------------------------------------------- URLs (R9 and Source URL cells only)
urls, seen = [], set()


def add(t):
    for u in BI.URL_RE.findall(t or ""):
        u = u.rstrip(".:")
        if "{" in u or "}" in u or (urlparse(u).hostname or "").endswith("googleapis.com"):
            continue
        if u not in seen:
            seen.add(u)
            urls.append(u)


for r in (20, 21):  # R9 Summary and Detail rows
    for c in range(PD0, LAST + 1):
        add(plain(ws.cell(row=r, column=c).value))
n9 = len(urls)
for b, (h, rows) in zip(blocks, pos):
    cc = b["hdr"].index("Source URL") + 1
    for r in rows:
        add(w3f.cell(row=r, column=cc).value)
open(URLS, "w", encoding="utf-8", newline="\n").write("\n".join(urls) + "\n")
print(f"urls: {len(urls)} unique ({n9} from AH20:AM21, {len(urls) - n9} new from 3f Source URL cells) -> {URLS}")

print("\nSUMMARY: %d/%d checks passed" % (sum(RESULTS), len(RESULTS)))
sys.exit(0 if all(RESULTS) else 1)
