"""Fill NeMo Guardrails columns (E..S) of sheet 3 from the approved md drafts."""
import re
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DIR = r"C:\Users\cys-c\Desktop\gzqr\benchtest"
XLSX = DIR + r"\AI Guardrails Research and Comparison.xlsx"
DRAFTS = [DIR + r"\drafts" + "\\" + n for n in
          ("jailbreak.md", "batch1.md", "batch2.md", "batch3.md")]
S3, S4 = "3. Guardrail Research Table", "4. Candidate Comparison Groups"

thin = Side(style="thin", color="8EA9DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
ALIGN = Alignment(wrap_text=True, vertical="top")

HDR_RE = re.compile(r"^## Column[^:]*:\s*((?:NeMo Guardrails|Llama Guard|GovTech Sentinel):.*)$")
CELL_RE = re.compile(r"^- \*\*R([1-9]):\*\* (.*)$")
LABEL_RE = re.compile(r"\[(Documented|Inferred|To be verified|Not disclosed)")


def parse_draft(path):
    cols, cur = [], None
    for line in open(path, encoding="utf-8").read().splitlines():
        if line.startswith("## Reviewer notes"):
            cur = None
            continue
        m = HDR_RE.match(line)
        if m:
            name = re.sub(r"\s*\(expanded\)\s*$", "", m.group(1).strip())
            cur = {"header": name, "cells": {}}
            cols.append(cur)
            continue
        if line.startswith("#"):
            continue
        m = CELL_RE.match(line)
        if m and cur is not None:
            n = int(m.group(1))
            assert n not in cur["cells"], (cur["header"], n)
            cur["cells"][n] = m.group(2).strip()
    for c in cols:
        assert sorted(c["cells"]) == list(range(1, 10)), (c["header"], sorted(c["cells"]))
    return cols


def to_runs(text, where):
    text = text.replace("`", "")
    parts = text.split("**")
    if len(parts) % 2 == 0:
        raise SystemExit(f"Unbalanced bold marker in {where}: {text[:80]}")
    return [(p, i % 2 == 1) for i, p in enumerate(parts) if p]


def build_cell(n, text, where):
    if n == 9:
        text = text.replace(" ; ", "\n")
    runs = to_runs(text, where)
    plain = "".join(t for t, _ in runs)
    if "**" in plain:
        raise SystemExit(f"Leftover ** in {where}")
    val = CellRichText(*[
        TextBlock(InlineFont(b=True, rFont="Arial", sz=10) if b
                  else InlineFont(rFont="Arial", sz=10), t) for t, b in runs])
    return val, runs, plain


def main():
    cols = []
    for p in DRAFTS:
        cols.extend(parse_draft(p))
    assert len(cols) == 15, len(cols)

    # approved text fix: Column M R4
    frag, repl = "All run on LLMRails only;", "The built-in flows run on LLMRails only;"
    mcol = [c for c in cols if c["header"] == "NeMo Guardrails: Grounded fact-checking (output)"]
    assert len(mcol) == 1
    assert mcol[0]["cells"][4].count(frag) == 1, "fragment not found exactly once"
    mcol[0]["cells"][4] = mcol[0]["cells"][4].replace(frag, repl)

    expected = {}
    for ci, c in enumerate(cols):
        for n in range(1, 10):
            expected[(n + 3, 5 + ci)] = build_cell(n, c["cells"][n], f"{c['header']} R{n}")

    before = load_workbook(XLSX, rich_text=True)
    wb = load_workbook(XLSX, rich_text=True)
    ws = wb[S3]

    # clear E onwards (old content / placeholders)
    for row in ws.iter_rows(min_row=3, min_col=5, max_col=max(ws.max_column, 19)):
        for c in row:
            c.value = None
            c.style = "Normal"
    assert 4 + len(cols) == 19
    for ci, c in enumerate(cols):
        x = ws.cell(row=3, column=5 + ci)
        x.value = c["header"]
        x.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", fgColor="4472C4")
        x.alignment, x.border = ALIGN, BORDER
    for (r, col), (val, _, _) in expected.items():
        x = ws.cell(row=r, column=col)
        x.value = val
        x.font = Font(name="Arial", size=10)
        if (r - 3) % 2 == 1:
            x.fill = PatternFill("solid", fgColor="DCE6F2")
        x.alignment, x.border = ALIGN, BORDER
    for col in range(5, 20):
        ws.column_dimensions[get_column_letter(col)].width = 48
    ws.freeze_panes = "D4"
    wb.save(XLSX)

    # ---- verification ----
    wb2 = load_workbook(XLSX, rich_text=True)
    ws = wb2[S3]
    print("Header A3:S3:")
    for c in ws[3][:19]:
        print("  ", c.coordinate, repr(c.value))
    print("max_column:", ws.max_column, "(must be 19)")
    print("freeze_panes:", ws.freeze_panes)
    fails, mism = 0, 0
    for r in range(4, 13):
        for col in range(5, 20):
            c = ws.cell(row=r, column=col)
            v = c.value
            blocks = list(v) if isinstance(v, CellRichText) else [v]
            plain = "".join(b.text if isinstance(b, TextBlock) else str(b) for b in blocks)
            ok = True
            first = blocks[0]
            if not (isinstance(first, TextBlock) and first.font.b):
                ok = False
                print("FAIL first-not-bold", c.coordinate)
            if r <= 10 and not any(isinstance(b, TextBlock) and b.font.b and LABEL_RE.search(b.text)
                                   for b in blocks):
                ok = False
                print("FAIL no bold label", c.coordinate)
            if "**" in plain or "`" in plain:
                ok = False
                print("FAIL marker", c.coordinate)
            if c.font.name != "Arial":
                ok = False
                print("FAIL font", c.coordinate)
            if "[Add Guardrail" in plain:
                ok = False
                print("FAIL placeholder", c.coordinate)
            fails += not ok
            if plain != expected[(r, col)][2]:
                mism += 1
                print("MISMATCH", c.coordinate)
    print("check failures:", fails)
    print("plain-text mismatches vs parsed drafts:", mism)

    def snap(w, cols_=None):
        out = []
        for row in w.iter_rows():
            for c in row:
                if cols_ and c.column not in cols_:
                    continue
                out.append((c.coordinate, str(c.value) if c.value is not None else None,
                            c.font.name, c.font.sz, c.font.b, c.fill.fgColor.rgb, c.fill.fill_type,
                            c.alignment.wrap_text,
                            c.border.left.color.rgb if c.border.left.color else None))
        return out
    print("columns A-D of sheet 3 unchanged:", snap(before[S3], {1, 2, 3, 4}) == snap(wb2[S3], {1, 2, 3, 4}))
    print("sheet 4 unchanged:", snap(before[S4]) == snap(wb2[S4])
          and before[S4].freeze_panes == wb2[S4].freeze_panes)
    mi = [i for i, c in enumerate(cols) if c["header"].endswith("Grounded fact-checking (output)")][0]
    print("Column M R4 (%s7):" % get_column_letter(5 + mi), str(ws.cell(row=7, column=5 + mi).value))
    print("S7 (Column N R4):", str(ws["S7"].value))
    print("Surprises scan:")
    for (r, col), (_, runs, _) in sorted(expected.items(), key=lambda k: (k[0][1], k[0][0])):
        ref = get_column_letter(col) + str(r)
        if not runs[0][1]:
            print("  non-bold first run:", ref)
        for t, b in runs:
            if not b and LABEL_RE.search(t):
                print("  non-bold label text:", ref, t[:60])


if __name__ == "__main__":
    main()
