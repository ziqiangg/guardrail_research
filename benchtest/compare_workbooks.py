"""Semantic comparison of two workbooks (values incl. formulas, rich-text runs, styles, merges, widths,
heights/outline, freeze, autofilter, data validation) for ALL sheets. openpyxl saves are not byte-reproducible
(zip timestamps), so use this instead of a byte diff.

Usage: python benchtest/compare_workbooks.py <a.xlsx> <b.xlsx>   (exit 0 = identical)
Tip:   git show HEAD:"benchtest/AI Guardrails Research and Comparison.xlsx" > /tmp/prev.xlsx
"""
import sys
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock

A, B = sys.argv[1], sys.argv[2]
wa, wb_ = load_workbook(A, rich_text=True), load_workbook(B, rich_text=True)
diffs = []


def d(*a):
    diffs.append(a)


def runs(v):
    if isinstance(v, CellRichText):
        out = []
        for b in v:
            if isinstance(b, TextBlock):
                f = b.font
                col = f.color.rgb if f is not None and f.color is not None else None
                out.append((b.text, f.b if f else None, f.i if f else None, f.sz if f else None, col,
                            f.rFont if f else None))
            else:
                out.append((str(b), None, None, None, None, None))
        return out
    return None


def plain(v):
    if isinstance(v, CellRichText):
        return "".join(b.text if isinstance(b, TextBlock) else str(b) for b in v)
    return v


if wa.sheetnames != wb_.sheetnames:
    d("sheet order", wa.sheetnames, wb_.sheetnames)
for n in wa.sheetnames:
    if n not in wb_.sheetnames:
        continue
    a, b = wa[n], wb_[n]
    if a.dimensions != b.dimensions:
        d(n, "dims", a.dimensions, b.dimensions)
    mr, mc = max(a.max_row, b.max_row), max(a.max_column, b.max_column)
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            x, y = a.cell(r, c), b.cell(r, c)
            k = (n, x.coordinate)
            if plain(x.value) != plain(y.value):
                d(*k, "value")
            if runs(x.value) != runs(y.value):
                d(*k, "runs")
            for attr in ("font", "fill", "border", "alignment", "number_format", "protection"):
                if repr(getattr(x, attr)) != repr(getattr(y, attr)):
                    d(*k, attr)
    if sorted(map(str, a.merged_cells.ranges)) != sorted(map(str, b.merged_cells.ranges)):
        d(n, "merges")
    ca = {k: (v.width, v.hidden, v.outline_level, v.min, v.max) for k, v in a.column_dimensions.items()}
    cb = {k: (v.width, v.hidden, v.outline_level, v.min, v.max) for k, v in b.column_dimensions.items()}
    if ca != cb:
        d(n, "col dims")
    ra = {k: (v.height, v.hidden, v.outline_level) for k, v in a.row_dimensions.items()}
    rb = {k: (v.height, v.hidden, v.outline_level) for k, v in b.row_dimensions.items()}
    if ra != rb:
        d(n, "row dims")
    if a.freeze_panes != b.freeze_panes:
        d(n, "freeze")
    if a.auto_filter.ref != b.auto_filter.ref:
        d(n, "autofilter")
    da = sorted((str(v.sqref), v.type, v.formula1, v.formula2) for v in a.data_validations.dataValidation)
    db = sorted((str(v.sqref), v.type, v.formula1, v.formula2) for v in b.data_validations.dataValidation)
    if da != db:
        d(n, "data validations")
    if a.sheet_properties.outlinePr != b.sheet_properties.outlinePr:
        d(n, "outlinePr")
print("sheets:", wa.sheetnames)
print("TOTAL DIFFS:", len(diffs))
for x in diffs[:30]:
    print(x)
sys.exit(1 if diffs else 0)
