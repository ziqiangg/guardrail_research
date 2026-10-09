"""Config-driven builder/verifier for the 'inventory' sheets (3d Llama Guard, 3e GovTech Sentinel).

A config is a dict with keys:
  md, sheet, title, note, after (sheet name to insert after), blocks [(md section marker, bold title or None,
  expected rows)], widths, center {block index: (1-based cols)}, covered (Covered-by header name),
  markers (strings that get the amber fill in the Covered-by column), order (expected sheet order),
  validators [callback(ctx) -> (ok, [messages])]  with ctx = dict(ws, ws3, cfg, blocks, pos, hdr3),
  panel (optional) coverage panel under the last block, see add_panel().
Public API: add_sheet(wb, cfg) -> ws ; verify(wb, cfg, xlsx=None)
Layout: title A1, source line A2, first table header on row 3 (autofilter on that table only), the other
tables each under a bold block title (one blank row between blocks); an md intro paragraph becomes a merged note.
"""
import math
import re
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import build_inventory as BI
from build_inventory import AMBER, BORDER, clean

GREY = "808080"


def cell_text(t):
    """Cleaned cell text; ' ; ' separated URL lists become one URL per line."""
    t = clean(t)
    if "http" in t and " ; " in t:
        t = t.replace(" ; ", "\n")
    return t


def parse_md(cfg):
    """Return [{title, intro, hdr, rows}] for the configured blocks; hdr/rows are cleaned cell texts."""
    lines = open(cfg["md"], encoding="utf-8").read().splitlines()
    blocks_cfg = cfg["blocks"]
    marks = [next(k for k, l in enumerate(lines) if l.startswith(m)) for m, _, _ in blocks_cfg]
    out = []
    for bi, (m, title, n) in enumerate(blocks_cfg):
        seg = lines[marks[bi] + 1:(marks[bi + 1] if bi + 1 < len(blocks_cfg) else len(lines))]
        intro = " ".join(l.strip() for l in seg if l.strip() and not l.startswith("|"))
        tbl = [[c.strip() for c in l.strip().strip("|").split("|")] for l in seg if l.startswith("|")]
        hdr, rows = tbl[0], tbl[2:]
        assert all(set(c) <= set("-: ") for c in tbl[1]), tbl[1]
        for r in rows:
            assert len(r) == len(hdr), (title, len(r), len(hdr), r[:2])
        assert len(rows) == n, (m, len(rows), n)
        out.append({"title": title, "intro": clean(intro), "hdr": [clean(h) for h in hdr],
                    "rows": [[cell_text(c) for c in r] for r in rows]})
    return out


def _body(ws, r, row, center=(), amber_cols=(), markers=()):
    for ci, v in enumerate(row, 1):
        c = ws.cell(row=r, column=ci, value=v)
        c.font = Font(name="Arial", size=10)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if ci in center else None)
        if ci in amber_cols and v.strip() in markers:
            c.fill = AMBER
    return r + 1


def _note(ws, r, text, W, widths):
    """Italic grey note merged across the block width; merged cells do not autofit, so set the row height."""
    c = ws.cell(row=r, column=1, value=text)
    c.font = Font(name="Arial", size=10, italic=True, color=GREY)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=W)
    chars = sum(widths[:W]) * 1.1
    ws.row_dimensions[r].height = 13.5 * max(1, math.ceil(len(text) / chars))


def add_sheet(wb, cfg):
    blocks = parse_md(cfg)
    name = cfg["sheet"]
    if name in wb.sheetnames:
        wb.remove(wb[name])
    ws = wb.create_sheet(name, wb.sheetnames.index(cfg["after"]) + 1)
    ws["A1"] = cfg["title"]
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    ws["A2"] = cfg["note"]
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color=GREY)
    r = 3
    where = []  # (header row, last data row) per block
    for bi, b in enumerate(blocks):
        W = len(b["hdr"])
        if b["title"]:
            r += 1  # blank row between blocks
            ws.cell(row=r, column=1, value=b["title"]).font = Font(name="Arial", size=10, bold=True)
            r += 1
        if b["intro"]:
            _note(ws, r, b["intro"], W, cfg["widths"])
            r += 1
        h = r
        BI._hdr(ws, r, b["hdr"])
        r += 1
        amber = (b["hdr"].index(cfg["covered"]) + 1,) if cfg["covered"] in b["hdr"] else ()
        for row in b["rows"]:
            r = _body(ws, r, row, center=cfg["center"].get(bi, ()), amber_cols=amber, markers=cfg["markers"])
        where.append((h, r - 1))
        if bi == 0:
            ws.auto_filter.ref = f"A{h}:{get_column_letter(W)}{r - 1}"
    for i, w in enumerate(cfg["widths"], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    if cfg.get("panel"):
        add_panel(ws, wb[BI.S3], cfg, blocks, [(h, list(range(h + 1, last + 1))) for h, last in where])
    return ws


# ------------------------------------------------------------------ coverage panel (below the last block)
def panel_cols(ws3, pnl):
    """Sheet-3 column indexes of the product (row-3 header prefix); checks the COUNTIF criteria are safe."""
    cols = [c for c in range(1, ws3.max_column + 1)
            if str(ws3.cell(row=3, column=c).value or "").startswith(pnl["prefix"])]
    assert len(cols) == pnl["n"] and cols == list(range(cols[0], cols[0] + pnl["n"])), cols
    hd = [ws3.cell(row=3, column=c).value for c in cols]
    for h in hd:
        assert not any(ch in h for ch in "*?~") and len(h) + 2 <= 255, h
        assert not any(h != o and h.lower() in o.lower() for o in hd), ("header is a substring of another", h)
    return cols, hd


def panel_range(blocks, pos, bi, header):
    """Absolute range of column `header` over the data rows of block bi, e.g. $O$4:$O$11."""
    L = get_column_letter(blocks[bi]["hdr"].index(header) + 1)
    rows = pos[bi][1]
    return f"${L}${rows[0]}:${L}${rows[-1]}"


def panel_cells(ws3, cfg, blocks, pos):
    """Return (start, cells {(row, col): value}, meta). start = title row = last block row + 3 (2 blank rows).
    cells hold the formulas exactly as written; meta lists the rows for the verifier."""
    pnl = cfg["panel"]
    cols, hd = panel_cols(ws3, pnl)
    cov = cfg["covered"]
    start = pos[-1][1][-1] + 3
    cells = {}
    nc = len(pnl["counts"])
    cells[(start, 1)] = pnl["title"]
    cells[(start + 1, 1)] = pnl["first_hdr"]
    for k, (lab, _bi) in enumerate(pnl["counts"]):
        cells[(start + 1, 2 + k)] = lab
    first = start + 2
    for i, c in enumerate(cols):
        r = first + i
        cells[(r, 1)] = f"='{BI.S3}'!{get_column_letter(c)}3"
        for k, (_lab, bi) in enumerate(pnl["counts"]):
            cells[(r, 2 + k)] = f'=COUNTIF({panel_range(blocks, pos, bi, cov)},"*"&$A{r}&"*")'
    r = first + len(cols)
    tot_rows = []
    for lab, terms in pnl["totals"]:
        cells[(r, 1)] = lab
        cells[(r, 2)] = "=" + "+".join(f'COUNTIF({panel_range(blocks, pos, bi, h)},"{crit}")'
                                       for bi, h, crit in terms)
        tot_rows.append(r)
        r += 1
    return start, cells, dict(nc=nc, first=first, ncols=len(cols), cols=cols, hd=hd, tot_rows=tot_rows, last=r - 1)


def panel_title(ws, r, text, last_col):
    """Merged panel title row (3b style); the merged-away cells also get the Arial font so every <c> is Arial."""
    BI._ptitle(ws, r, text, last_col)
    for c in range(2, last_col + 1):
        ws.cell(row=r, column=c).font = Font(name="Arial", size=10, bold=True, color="FFFFFF")


def add_panel(ws, ws3, cfg, blocks, pos):
    start, cells, meta = panel_cells(ws3, cfg, blocks, pos)
    nc = meta["nc"]
    panel_title(ws, start, cells[(start, 1)], 1 + nc)
    for c in range(1, 2 + nc):
        BI._pcell(ws, start + 1, c, cells[(start + 1, c)], hdr=True)
    for (r, c), v in cells.items():
        if r < start + 2:
            continue
        BI._pcell(ws, r, c, v, num=(c >= 2), bold=(r in meta["tot_rows"] and c == 1))
    for r in meta["tot_rows"]:  # empty bordered cells next to the single total
        for c in range(3, 2 + nc):
            BI._pcell(ws, r, c, None, num=True)


def _crit_re(crit):
    """COUNTIF text criterion (wildcards * only; case-insensitive; * also spans line breaks) as a regex."""
    assert "?" not in crit and "~" not in crit
    return re.compile("".join(".*" if ch == "*" else re.escape(ch) for ch in crit), re.I | re.S)


def python_panel(cfg, ws, ws3, blocks, pos):
    """Python mirror of every panel value: ({header: [count per count column]}, {total label: value})."""
    pnl, cov = cfg["panel"], cfg["covered"]
    cols, hd = panel_cols(ws3, pnl)

    def col_vals(bi, header):
        ci = blocks[bi]["hdr"].index(header) + 1
        return [str(ws.cell(row=r, column=ci).value) for r in pos[bi][1]]
    counts = {h: [sum(bool(_crit_re("*" + h + "*").fullmatch(v)) for v in col_vals(bi, cov))
                  for _lab, bi in pnl["counts"]] for h in hd}
    totals = {lab: sum(bool(_crit_re(crit).fullmatch(v)) for bi, h, crit in terms for v in col_vals(bi, h))
              for lab, terms in pnl["totals"]}
    return counts, totals


def verify_panel(wb, cfg, ws, ws3, blocks, pos):
    """Panel cells equal the specified formulas; Python mirror values printed; every product column covered."""
    pnl = cfg["panel"]
    start, cells, meta = panel_cells(ws3, cfg, blocks, pos)
    got = {(c.row, c.column): c.value for row in ws.iter_rows(min_row=pos[-1][1][-1] + 1) for c in row
           if c.value is not None}
    ok = got == cells
    print("panel: title row %d (last block row %d + 3 = 2 blank rows), %d cells, equal to the specified formulas: %s"
          % (start, pos[-1][1][-1], len(cells), "OK" if ok else "FAIL"))
    if not ok:
        print("  diff:", [(k, got.get(k), cells.get(k)) for k in sorted(set(got) | set(cells))
                          if got.get(k) != cells.get(k)][:5])
    blank = all(c.value is None for r in (start - 2, start - 1) for c in ws[r])
    fonts = [c.coordinate for row in ws.iter_rows(min_row=start) for c in row
             if type(c).__name__ != "MergedCell" and c.value is not None and c.font.name != "Arial"]
    print("two rows above the title empty:", blank, "| non-Arial panel cells:", len(fonts))
    counts, totals = python_panel(cfg, ws, ws3, blocks, pos)
    print("Python mirror counts per Table 3 column", [lab for lab, _ in pnl["counts"]])
    for h, v in counts.items():
        print("   ", v, h)
    for lab, v in totals.items():
        print("    total %-52s %d" % (lab, v))
    unc = [h for h, v in counts.items() if v[-1] == 0]  # last count column: variants (3d) / guardrail IDs (3e)
    print("Table 3 columns with a zero count:", unc)
    ok = ok and blank and not fonts and not unc
    return ok, counts, totals


# ------------------------------------------------------------------ verification
def block_rows(ws, first_hdr_col_a):
    """(header row, data rows) of the table whose header has column A == first_hdr_col_a (rows until a blank A)."""
    h = next(r for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value == first_hdr_col_a)
    rows, k = [], h + 1
    while k <= ws.max_row and ws.cell(row=k, column=1).value:
        rows.append(k)
        k += 1
    return h, rows


def verify(wb, cfg, xlsx=None):
    from build_eval_sheet import xml_non_arial
    name = cfg["sheet"]
    print("\n=== inventory_sheet verification (%s) ===" % name)
    exp_order = cfg["order"]
    print("sheet order:", wb.sheetnames, "OK" if wb.sheetnames == exp_order else "FAIL (want %s)" % exp_order)
    ws, ws3 = wb[name], wb[BI.S3]
    blocks = parse_md(cfg)
    nw = len(cfg["widths"])
    print("A1:", ws["A1"].value, "| bold/size:", ws["A1"].font.b, ws["A1"].font.sz,
          "| A2 italic/color:", ws["A2"].font.i, ws["A2"].font.color.rgb)
    print("freeze:", ws.freeze_panes, "(A4)  max_row:", ws.max_row, " widths:",
          [ws.column_dimensions[get_column_letter(i)].width for i in range(1, nw + 1)])

    mism = total = 0
    pos = []
    for b in blocks:
        h, rows = block_rows(ws, b["hdr"][0])
        pos.append((h, rows))
        got = [[ws.cell(row=r, column=c).value for c in range(1, len(b["hdr"]) + 1)] for r in [h] + rows]
        want = [b["hdr"]] + b["rows"]
        total += sum(len(x) for x in want)
        nm = sum(1 for g, w in zip(got, want) for x, y in zip(g, w) if x != y) + abs(len(got) - len(want))
        mism += nm
        extra = [c.coordinate for r in [h] + rows for c in ws[r][len(b["hdr"]):] if c.value is not None]
        print(f"  block '{b['hdr'][0]}': header row {h}, rows {len(rows)} (md {len(b['rows'])}),",
              "mismatches", nm, "| cells beyond block width:", len(extra))
        mism += len(extra)
    print("cell text vs cleaned md (%d cells compared): %d mismatches" % (total, mism))

    counts_ok = all(len(p[1]) == n for p, (_, _, n) in zip(pos, cfg["blocks"]))
    h1, r1 = pos[0]
    W0 = len(blocks[0]["hdr"])
    H1_WANT = 3 + (1 if blocks[0]["intro"] else 0)  # an md intro paragraph of block (a) becomes a note row above it
    want_af = f"A{h1}:{get_column_letter(W0)}{r1[-1]}"
    print("block row counts:", [len(p[1]) for p in pos], "OK" if counts_ok else "FAIL", "| first header row", h1,
          "(%d)" % H1_WANT, "OK" if h1 == H1_WANT else "FAIL")
    print("autofilter:", ws.auto_filter.ref, "expected", want_af, "OK" if ws.auto_filter.ref == want_af else "FAIL")

    # covered-by (every block that has the column)
    hdr3 = {ws3.cell(row=3, column=c).value for c in range(1, ws3.max_column + 1)}
    markers = set(cfg["markers"])
    bad, n_mark, amber_ok = [], 0, True
    for b, (h, rows) in zip(blocks, pos):
        if cfg["covered"] not in b["hdr"]:
            continue
        cc = b["hdr"].index(cfg["covered"]) + 1
        for r in rows:
            for v in str(ws.cell(row=r, column=cc).value).split(";"):
                v = v.strip()
                n_mark += v in markers
                if v not in markers and v not in hdr3:
                    bad.append((r, v))
        amb = [r for r in rows if (ws.cell(row=r, column=cc).fill.fgColor.rgb or "").endswith("FFF2CC")]
        mk_rows = [r for r in rows if ws.cell(row=r, column=cc).value.strip() in markers]
        amber_ok = amber_ok and amb == mk_rows
    print("covered-by values not a sheet-3 row-3 header or marker:", len(bad), bad[:3],
          "| marker cells:", n_mark, "| amber cells == marker cells:", amber_ok)

    ctx = dict(ws=ws, ws3=ws3, cfg=cfg, blocks=blocks, pos=pos, hdr3=hdr3)
    vok = True
    if cfg.get("panel"):
        pok, _, _ = verify_panel(wb, cfg, ws, ws3, blocks, pos)
        vok = vok and pok
    for fn in cfg.get("validators", []):
        ok_, msgs = fn(ctx)
        for m in msgs:
            print("  " + m)
        vok = vok and ok_

    marks = [c.coordinate for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and ("**" in c.value or "`" in c.value)]
    print("cells with ** or backtick on %s: %d %s" % (name, len(marks), marks[:5]))
    nonar = [c.coordinate for row in ws.iter_rows() for c in row
             if type(c).__name__ != "MergedCell" and c.value is not None and c.font.name != "Arial"]
    print("non-Arial cells with values:", len(nonar), nonar[:5])
    xbad = 0
    if xlsx:
        n, xbad = xml_non_arial(xlsx, name)
        print("raw-XML font check (all %d <c> elements): non-Arial = %d" % (n, xbad))
    ok = (not mism and not bad and amber_ok and not marks and not nonar and not xbad and vok and counts_ok
          and h1 == H1_WANT and ws.auto_filter.ref == want_af and wb.sheetnames == exp_order)
    print("%s overall:" % name, "PASS" if ok else "FAIL")
    assert ok
    return mism
