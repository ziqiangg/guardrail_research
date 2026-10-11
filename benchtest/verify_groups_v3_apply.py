"""Independent verification of the sheet-4 regroup (groups_v3.md, 34 rows) against the previous commit's workbook
(git show HEAD:...).  Every other sheet must be unchanged.  Exit 0 when every check passes.
"""
import os
import re
import subprocess
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter

import build_two_level as B2
import build_inventory as BI
from build_eval_sheet import xml_non_arial

XLSX = B2.XLSX
S3, S4 = BI.S3, BI.S4
MD = str(BI.P.DRAFTS / "groups_v3.md")
RESULTS = []


def check(name, ok, extra=""):
    RESULTS.append(bool(ok))
    print(("PASS" if ok else "FAIL"), "-", name, extra)


plain = B2.plain_of


def runs_of(v):
    if isinstance(v, CellRichText):
        return [(b.text, bool(b.font.b), b.font.sz, b.font.rFont) for b in v]
    return [(plain(v), False, None, None)]


def _col(c):
    return None if c is None else (c.rgb if c.type == "rgb" else (c.type, c.theme if c.type == "theme" else c.indexed))


def style_of(c):
    b = c.border
    return (c.font.name, c.font.sz, c.font.b, c.font.i, _col(c.font.color), c.fill.fill_type, _col(c.fill.fgColor),
            c.alignment.wrap_text, c.alignment.vertical, c.alignment.horizontal,
            tuple((s.style, _col(s.color)) if s is not None else None for s in (b.left, b.right, b.top, b.bottom)))


def full_snap(w):
    cells = [(c.coordinate, plain(c.value), runs_of(c.value), style_of(c)) for row in w.iter_rows() for c in row]
    layout = (sorted(str(m) for m in w.merged_cells.ranges), w.freeze_panes, w.auto_filter.ref,
              {k: v.width for k, v in w.column_dimensions.items()},
              {k: (v.height, v.outline_level, v.hidden) for k, v in w.row_dimensions.items()},
              (w.max_row, w.max_column),
              w.sheet_properties.outlinePr.summaryBelow if w.sheet_properties.outlinePr else None)
    return cells, layout


prev = os.path.join(tempfile.mkdtemp(), "prev.xlsx")
with open(prev, "wb") as f:
    f.write(subprocess.run(["git", "show", 'HEAD:benchtest/AI Guardrails Research and Comparison.xlsx'],
                           capture_output=True, check=True, cwd=str(BI.P.ROOT)).stdout)
new = load_workbook(XLSX, rich_text=True)
old = load_workbook(prev, rich_text=True)
check("sheet names and order identical to the previous commit, sheet 4 last",
      new.sheetnames == old.sheetnames == list(BI.ORDER) and new.sheetnames[-1] == S4, str(len(new.sheetnames)))
for nm in old.sheetnames:
    if nm == S4:
        continue
    a, al = full_snap(new[nm])
    b, bl = full_snap(old[nm])
    check(f"{nm}: values, rich text, styles, merges, widths, freeze, filter, rows identical", a == b and al == bl,
          f"({sum(x != y for x, y in zip(a, b))} cell diffs, {len(a)} vs {len(b)} cells)")

# ---------------------------------------------------------------- groups_v3.md, parsed here from scratch
lines = open(MD, encoding="utf-8").read().splitlines()
i0 = next(k for k, l in enumerate(lines) if l.startswith("## A."))
rows = []
for l in lines[i0 + 1:]:
    if l.startswith("## B."):
        break
    if l.startswith("|"):
        rows.append([c.strip() for c in l.strip().strip("|").split("|")])
hdr, rows = rows[0], rows[2:]
check("md table: 8 columns, 34 rows, C1..C34 in order",
      len(hdr) == 8 and len(rows) == 34 and all(r[0].startswith(f"C{k + 1} ") and len(r) == 8
                                               for k, r in enumerate(rows)), f"({len(rows)} rows)")
ws, ws3 = new[S4], new[S3]
check("A1 title and row-3 header equal the md header / previous commit (values and styles)",
      [ws.cell(row=3, column=c).value for c in range(1, 9)] == hdr and ws["A1"].value == old[S4]["A1"].value
      and [style_of(ws.cell(row=3, column=c)) for c in range(1, 9)] ==
      [style_of(old[S4].cell(row=3, column=c)) for c in range(1, 9)])
grows = list(range(5, 19)) + list(range(20, 40))
check("bands at rows 4 and 19 with the expected text, A:H grey D9D9D9, bold Arial, unmerged",
      ws["A4"].value == "Comparison groups (2+ products)"
      and ws["A19"].value == "Single-product functions (no comparator yet)"
      and all(ws.cell(row=r, column=c).fill.fgColor.rgb.endswith("D9D9D9") and ws.cell(row=r, column=c).font.b
              and ws.cell(row=r, column=c).font.name == "Arial" for r in (4, 19) for c in range(1, 9))
      and not [m for m in ws.merged_cells.ranges if m.min_row < 40])
check("freeze B4, autofilter A3:H39, widths 28/40/42x6",
      ws.freeze_panes == "B4" and ws.auto_filter.ref == "A3:H39" and
      [ws.column_dimensions[get_column_letter(c)].width for c in range(1, 9)] == [28, 40] + [42] * 6)

mism = 0
nbul = nref = 0
bad = []
for k, r in enumerate(rows):
    row = grows[k]
    single = k >= 14
    marker = r[7].startswith("• Single-product: no comparator among the ten products yet")
    if marker != single:
        bad.append(("marker", k + 1))
    for c, v in enumerate(r, 1):
        cell = ws.cell(row=row, column=c)
        want = "\n".join(p.strip() for p in v.split("; ")) if c == 2 else v.replace("<br>", "\n")
        if plain(cell.value) != want:
            mism += 1
        if c >= 3:
            ls = v.split("<br>")
            if not ls[-1].startswith("Refs: "):
                bad.append(("no refs", k + 1, c))
            bl = list(cell.value)
            if len(bl) != len(ls):
                bad.append(("runs", row, c))
                continue
            for b, ln in zip(bl, ls):
                if ln.startswith("Refs: "):
                    nref += 1
                    if not (b.font.sz == 8 and b.font.color.rgb.endswith("808080") and b.font.rFont == "Arial"):
                        bad.append(("refs style", row, c))
                else:
                    nbul += 1
                    w = len(ln[2:].split())
                    if not ln.startswith("• ") or w > 12:
                        bad.append(("bullet", row, c, w))
                    if b.font.rFont != "Arial" or b.font.sz != 10:
                        bad.append(("bullet font", row, c))
check("34 rows x 8 cells: plain text equals groups_v3.md (rows 5-18, 20-39)", mism == 0, f"({mism} mismatches)")
check("single-product marker exactly on C15-C34 and nowhere on C1-C14; each cell ends with a grey Refs line "
      "(Arial 8 808080); bullets start with a bullet char, at most 12 words, Arial 10", not bad,
      f"({nbul} bullets, {nref} Refs lines; bad: {bad[:4]})")
check("all sheet 4 cells Arial; nothing beyond column H above the panel",
      all(c.font.name == "Arial" for row in ws.iter_rows() for c in row
          if type(c).__name__ != "MergedCell" and c.value is not None)
      and not [1 for row in ws.iter_rows(max_row=39) for c in row if c.column > 8 and c.value])
nn, bd = xml_non_arial(XLSX, S4)
check("raw XML: every sheet-4 <c> Arial", bd == 0, f"({bd} of {nn})")

# ---------------------------------------------------------------- coverage panel
hd = [ws3.cell(row=3, column=c).value for c in range(5, 67)]
check("sheet 3 has 62 function columns E..BN, headers unique and none a substring of another",
      len(hd) == 62 and all(isinstance(h, str) for h in hd) and
      not any(h != o and h.lower() in o.lower() for h in hd for o in hd))
bcol = [plain(ws.cell(row=r, column=2).value).lower() for r in grows]
cnt = [sum(h.lower() in b for b in bcol) for h in hd]
letters = {}
for k, r in enumerate(rows):
    for ln in r[1].split("; "):
        m = re.match(r"^([A-Z]{1,2}): (.*)$", ln.strip())
        letters.setdefault(m.group(1), set()).add(k)
        col = next(c for c in range(1, 80) if get_column_letter(c) == m.group(1))
        assert ln.strip().startswith(f"{m.group(1)}: {hd[col - 5]}"), ln
check("every function line is '<letter>: <sheet 3 header>'; per-function mirror counts equal the lines per letter",
      cnt == [len(letters.get(get_column_letter(5 + i), ())) for i in range(62)])
first = 43
exp = {}
for i, h in enumerate(hd):
    r = first + i
    exp[(r, 1)] = get_column_letter(5 + i)
    exp[(r, 2)] = f"='{S3}'!{get_column_letter(5 + i)}3"
    exp[(r, 3)] = f'=COUNTIF($B$5:$B$39,"*"&$B{r}&"*")'
cc = f"$C${first}:$C${first + 61}"
tot = [("Functions in no group", f'=COUNTIF({cc},0)'), ("Functions in more than one group", f'=COUNTIF({cc},">1")')]
prod = [("NeMo Guardrails", ["NeMo Guardrails:"]), ("Llama Guard", ["Llama Guard:"]),
        ("GovTech Sentinel", ["GovTech Sentinel:"]), ("Presidio", ["Presidio:"]),
        ("Sensitive Data Protection", ["Sensitive Data Protection:"]), ("Model Armor", ["Model Armor:"]),
        ("LionGuard", ["LionGuard:"]), ("Purple Llama", ["Prompt Guard 2:", "LlamaFirewall:", "Code Shield:"]),
        ("Cloak", ["Cloak:"]), ("Amazon Bedrock", ["Amazon Bedrock"])]
mir = {}
for nm, pats in prod:
    if len(pats) == 1:
        f = f'=COUNTIF($B$5:$B$39,"*{pats[0]}*")'
    else:
        f = "=SUMPRODUCT(--((" + "+".join(f'ISNUMBER(SEARCH("{q}",$B$5:$B$39))' for q in pats) + ")>0))"
    tot.append((f"Groups including {nm}", f))
    mir[nm] = sum(any(q.lower() in b for q in pats) for b in bcol)
tot.append(("Single-product rows", '=COUNTIF($H$5:$H$39,"*Single-product:*")'))
for k, (lab, f) in enumerate(tot):
    exp[(first + 62 + k, 1)] = lab
    exp[(first + 62 + k, 3)] = f
exp[(41, 1)] = "Coverage of sheet 3 functions"
for k, h in enumerate(["Sheet 3 column", "Function (live link)", "Groups listing it"]):
    exp[(42, 1 + k)] = h
got = {(c.row, c.column): c.value for row in ws.iter_rows(min_row=40) for c in row if c.value is not None}
check("panel (title row 41, header 42, 62 function rows 43-104, 13 totals 105-117) equals the independently "
      "specified formulas", got == exp,
      "" if got == exp else str([(k, got.get(k), exp.get(k)) for k in sorted(set(got) | set(exp))
                                 if got.get(k) != exp.get(k)][:3]))
check("panel title merged A41:C41; row 40 blank; panel cells Arial; no dynamic-array functions",
      "A41:C41" in {str(m) for m in ws.merged_cells.ranges} and all(c.value is None for c in ws[40]) and
      all(c.font.name == "Arial" for row in ws.iter_rows(min_row=40) for c in row
          if type(c).__name__ != "MergedCell" and c.value is not None) and
      not any(re.search(r"FILTER|UNIQUE|XLOOKUP|LET\(|SORT\(|SEQUENCE", str(v)) for v in got.values()))
multi = sum(c > 1 for c in cnt)
print("  Python mirror: in no group", cnt.count(0), "| more than one group", multi, "|", mir,
      "| single-product rows", sum(r[7].startswith("• Single-product:") for r in rows))
check("0 functions in no group; 23 in more than one group (groups_v3.md accounting)",
      cnt.count(0) == 0 and multi == 23)
check("md accounting line (62 functions, 0 in no group, 23 in more than one group, 34 groups: 14 + 20) present",
      any("Totals: 62 functions (E to BN); 0 in no group; 23 in more than one group; 34 groups "
          "(14 multi-product, 20 single-product)" in l for l in lines))

# ---------------------------------------------------------------- XML hygiene
ssp, emp, norpr, n, wf = BI.hygiene(XLSX)
check("XML hygiene (all parts well-formed)", ssp == 0 and emp == 0 and norpr == 0 and wf,
      f"(unpreserved <t>={ssp}, <r> without rPr={norpr}, empty <t>={emp}, parts={n}, well-formed={wf})")

print("\n%d/%d checks passed" % (sum(RESULTS), len(RESULTS)))
sys.exit(0 if all(RESULTS) else 1)
