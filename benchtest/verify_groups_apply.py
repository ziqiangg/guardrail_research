"""Independent verification of the sheet-4 readability update (bulleted groups_v2.md, bands, panel below) against .v7.xlsx.

NOTE (2026-10-11): pinned to the v7 baseline and the 24-row groups_v2 layout; superseded by verify_groups_v3_apply.py. Do not run against the current workbook.
Run: python verify_groups_apply.py     (exit code 0 when every check passes)
Expected formulas are spelled out here from scratch (not taken from the builders) so the check is independent.
"""
import re, sys
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils import get_column_letter

import build_two_level as B2
import build_inventory as BI
import build_groups_sheet as BG
from build_eval_sheet import xml_non_arial

XLSX = B2.XLSX
V7 = B2.BAK7
S3, S3B, S3C, S3D, S3E, S4 = BI.S3, BI.S3B, BI.S3C, BI.S3D, BI.S3E, BI.S4
RESULTS = []


def check(name, ok, extra=""):
    RESULTS.append(bool(ok))
    print(("PASS" if ok else "FAIL"), "-", name, extra)


plain = B2.plain_of


def _col(c):
    return None if c is None else (c.rgb if c.type == "rgb" else (c.type, c.theme if c.type == "theme" else c.indexed))


def style_of(c):
    b = c.border
    return (c.font.name, c.font.sz, c.font.b, c.font.i, _col(c.font.color), c.fill.fill_type, _col(c.fill.fgColor),
            c.alignment.wrap_text, c.alignment.vertical, c.alignment.horizontal,
            tuple((s.style, _col(s.color)) if s is not None else None for s in (b.left, b.right, b.top, b.bottom)))


def snap(w, max_row=None):
    """Values (incl. formula strings), styles, merges, layout; max_row limits to the original rows."""
    top = max_row or w.max_row
    cells = [(c.coordinate, plain(c.value), style_of(c)) for row in w.iter_rows(max_row=top) for c in row]
    layout = (sorted(str(m) for m in w.merged_cells.ranges if m.min_row <= top), w.freeze_panes, w.auto_filter.ref,
              {k: v.width for k, v in w.column_dimensions.items()},
              {k: (v.height, v.outline_level, v.hidden) for k, v in w.row_dimensions.items() if k <= top},
              (top, w.max_column) if max_row else (w.max_row, w.max_column))
    return cells, layout


new = load_workbook(XLSX, rich_text=True)
old = load_workbook(V7, rich_text=True)
print("v7 sheets:", old.sheetnames)
check("sheet order unchanged (6 sheets, sheet 4 last)", new.sheetnames == old.sheetnames ==
      BI.ORDER and len(new.sheetnames) == len(BI.ORDER) and new.sheetnames[-1] == S4)

# ---------------------------------------------------------------- sheets 3, 3b, 3c, 3d, 3e identical to v7 (3d/3e incl. panels)
for nm in (S3, S3B, S3C, S3D, S3E):
    a, al = snap(new[nm])
    b, bl = snap(old[nm])
    nd = sum(1 for x, y in zip(a, b) if x != y) + abs(len(a) - len(b))
    check(f"{nm}: values (incl. formulas), fonts, fills, borders, alignment identical to v7", a == b,
          f"({len(a)} cells, {nd} diffs)")
    check(f"{nm}: merges, freeze, autofilter, widths, row heights/outline, extents identical to v7", al == bl,
          "" if al == bl else str([(i, x, y) for i, (x, y) in enumerate(zip(al, bl)) if x != y])[:300])

ws3 = new[S3]
hdr3 = {c: ws3.cell(row=3, column=c).value for c in range(1, 34)}

# ---------------------------------------------------------------- 3d / 3e
LGC, SNC = "Llama Guard:", "GovTech Sentinel:"
COVH = "Covered by Table 3 column"


def cols_of(prefix, n):
    cs = [c for c, h in hdr3.items() if str(h).startswith(prefix)]
    assert len(cs) == n, cs
    return cs


def crit_re(crit):
    return re.compile("".join(".*" if ch == "*" else re.escape(ch) for ch in crit), re.I | re.S)


def mirror(vals, crit):
    r = crit_re(crit)
    return sum(bool(r.fullmatch(str(v))) for v in vals)


def find_rows(ws, first_hdr, hdr_name):
    """(header row, column letter of hdr_name, data rows) of the table whose header row has A == first_hdr."""
    h = next(r for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value == first_hdr)
    col = next(c for c in range(1, 17) if ws.cell(row=h, column=c).value == hdr_name)
    rows, k = [], h + 1
    while ws.cell(row=k, column=1).value:
        rows.append(k)
        k += 1
    return h, col, rows


def check_product(name, prefix, n, blocks, counts, totals):
    """blocks: {key: (first header, covered/other header)}; counts: [(label, block key)]; totals: [(label, [(key, header, crit)])]."""
    w, wo = new[name], old[name]
    hp0 = next(r for r in range(3, wo.max_row + 1) if wo.cell(row=r, column=1).value == "Path")
    top = hp0
    while wo.cell(row=top + 1, column=1).value:
        top += 1  # last data row of the original tables (the v7 sheet also holds the panel below)
    a, al = snap(w, top)
    b, bl = snap(wo, top)
    nd = sum(1 for x, y in zip(a, b) if x != y) + abs(len(a) - len(b))
    check(f"{name}: all {top} original rows identical to v7 (values + styles, {len(a)} cells)", a == b, f"({nd} diffs)")
    check(f"{name}: original merges, freeze, autofilter ({wo.auto_filter.ref}), widths, row heights identical to v7",
          al == bl and w.auto_filter.ref == wo.auto_filter.ref, "" if al == bl else str(
              [(i, x, y) for i, (x, y) in enumerate(zip(al, bl)) if x != y])[:300])
    extra = sorted(str(m) for m in w.merged_cells.ranges if m.min_row > top)
    check(f"{name}: only new merge = the panel title row", len(extra) == 1 and extra[0].startswith("A"), str(extra))
    # block positions
    pos = {k: find_rows(w, fh, ch) for k, (fh, ch) in blocks.items()}
    hp = next(r for r in range(3, w.max_row + 1) if w.cell(row=r, column=1).value == "Path")  # last block: 'Path'
    last_row = hp
    while w.cell(row=last_row + 1, column=1).value:
        last_row += 1
    check(f"{name}: last block ('Path' table) ends at row {last_row} = original max_row {top}", last_row == top)
    # panel location: 2 blank rows below the last block
    start = last_row + 3
    check(f"{name}: panel title at row {start} = last block row {last_row} + 3 (2 blank rows)",
          w.cell(row=start, column=1).value == "Coverage of Table 3 columns" and
          all(c.value is None for r in (last_row + 1, last_row + 2) for c in w[r]))
    cols = cols_of(prefix, n)
    want = {}
    want[(start, 1)] = "Coverage of Table 3 columns"
    want[(start + 1, 1)] = "Table 3 column (live link)"
    for k, (lab, _) in enumerate(counts):
        want[(start + 1, 2 + k)] = lab
    rng = lambda key, header: (lambda col, rows: f"${get_column_letter(col)}${rows[0]}:${get_column_letter(col)}${rows[-1]}")(
        next(c for c in range(1, 17) if w.cell(row=pos[key][0], column=c).value == header), pos[key][2])
    for i, c in enumerate(cols):
        r = start + 2 + i
        want[(r, 1)] = f"='{S3}'!{get_column_letter(c)}3"
        for k, (lab, key) in enumerate(counts):
            want[(r, 2 + k)] = f'=COUNTIF({rng(key, COVH)},"*"&$A{r}&"*")'
    r = start + 2 + n
    mirrors = {}
    for lab, terms in totals:
        want[(r, 1)] = lab
        want[(r, 2)] = "=" + "+".join(f'COUNTIF({rng(key, hh)},"{cr}")' for key, hh, cr in terms)
        vals = 0
        for key, hh, cr in terms:
            ci = next(c for c in range(1, 17) if w.cell(row=pos[key][0], column=c).value == hh)
            vals += mirror([w.cell(row=x, column=ci).value for x in pos[key][2]], cr)
        mirrors[lab] = vals
        r += 1
    got = {(c.row, c.column): c.value for row in w.iter_rows(min_row=top + 1) for c in row if c.value is not None}
    check(f"{name}: appended rows hold exactly the panel ({len(want)} cells) with the specified formulas", got == want,
          "" if got == want else str([(k, got.get(k), want.get(k)) for k in sorted(set(got) | set(want))
                                      if got.get(k) != want.get(k)][:4]))
    check(f"{name}: no content between the original rows and the panel", not any(
        c.value is not None for row in w.iter_rows(min_row=top + 1, max_row=start - 1) for c in row))
    hds = [hdr3[c] for c in cols]
    check(f"{name}: panel links resolve to the {n} '{prefix}' headers V..", [ws3.cell(row=3, column=c).value for c in cols] == hds)
    check(f"{name}: header guards (no *?~, <=253 chars, none a substring of another)",
          all(not any(ch in h for ch in "*?~") and len(h) + 2 <= 255 for h in hds) and
          not any(h != o and h.lower() in o.lower() for h in hds for o in hds))
    counts_mirror = {}
    for h in hds:
        counts_mirror[h] = []
        for lab, key in counts:
            ci = next(c for c in range(1, 17) if w.cell(row=pos[key][0], column=c).value == COVH)
            counts_mirror[h].append(mirror([w.cell(row=x, column=ci).value for x in pos[key][2]], "*" + h + "*"))
    print("   Python mirror counts", [lab for lab, _ in counts])
    for h, v in counts_mirror.items():
        print("     ", v, h)
    for lab, v in mirrors.items():
        print("      total %-62s %d" % (lab, v))
    check(f"{name}: every Table 3 column covered by >=1 {counts[-1][0].split(' covering')[0]}",
          all(v[-1] >= 1 for v in counts_mirror.values()), "" if all(v[-1] >= 1 for v in counts_mirror.values()) else
          str([h for h, v in counts_mirror.items() if v[-1] < 1]))
    nonar = [c.coordinate for row in w.iter_rows(min_row=top + 1) for c in row
             if type(c).__name__ != "MergedCell" and c.value is not None and c.font.name != "Arial"]
    nx, bx = xml_non_arial(XLSX, name)
    check(f"{name}: panel + sheet all Arial (cells and raw XML)", not nonar and bx == 0, f"(cells {len(nonar)}, raw {bx}/{nx})")
    mk = [c.coordinate for row in w.iter_rows() for c in row if isinstance(c.value, str) and ("**" in c.value or "`" in c.value)]
    check(f"{name}: no ** or backticks", not mk)
    return counts_mirror, mirrors


lg_cm, lg_t = check_product(S3D, LGC, 5, {"a": ("Variant", COVH)}, [("Variants covering it", "a")],
                            [("Variants marked legacy", [("a", COVH, "*legacy*")])])
STAT, OWN, IO = "Status (Available/Planned)", "Owner (GovTech/AWS/Meta)", "Input/Output"
sn_cm, sn_t = check_product(
    S3E, SNC, 7, {"a": ("Variant", COVH), "b": ("Guardrail ID", COVH)},
    [("Model variants covering it", "a"), ("Guardrail IDs covering it", "b")],
    [("Guardrail IDs Available", [("b", STAT, "Available*")]),
     ("Guardrail IDs Planned", [("b", STAT, "Planned*")]),
     ("Guardrail IDs owned by GovTech", [("b", OWN, "GovTech*")]),
     ("Guardrail IDs owned by AWS", [("b", OWN, "AWS*")]),
     ("Guardrail IDs owned by Meta (model; listed under owner govtech)", [("b", OWN, "Listed under owner govtech*")]),
     ("Guardrail IDs: Input/Output (first stated)", [("b", IO, "Input/Output*")]),
     ("Guardrail IDs: Input only (first stated)", [("b", IO, "Input *")]),
     ("Guardrail IDs: Output only (first stated)", [("b", IO, "Output*")]),
     ("Covered-by values marked planned", [("a", COVH, "*planned*"), ("b", COVH, "*planned*")]),
     ("Covered-by values marked legacy", [("a", COVH, "*legacy*"), ("b", COVH, "*legacy*")])])
check("3e panel partitions: status 22+2=24, owners 16+7+1=24, input/output 13+8+3=24",
      sn_t["Guardrail IDs Available"] + sn_t["Guardrail IDs Planned"] == 24 and
      sum(v for k, v in sn_t.items() if k.startswith("Guardrail IDs owned")) == 24 and
      sum(v for k, v in sn_t.items() if k.startswith("Guardrail IDs: ")) == 24)
check("3d/3e expected totals: legacy variants 2; Available 22, Planned 2; planned values 2; legacy values 1",
      lg_t["Variants marked legacy"] == 2 and sn_t["Guardrail IDs Available"] == 22 and sn_t["Guardrail IDs Planned"] == 2
      and sn_t["Covered-by values marked planned"] == 2 and sn_t["Covered-by values marked legacy"] == 1)

# ---------------------------------------------------------------- sheet 4
w4, o4 = new[S4], old[S4]
check("sheet 4: A1 and header row 3 (values, font, fill, border, alignment) identical to v7",
      all(plain(w4.cell(row=r, column=c).value) == plain(o4.cell(row=r, column=c).value) and
          style_of(w4.cell(row=r, column=c)) == style_of(o4.cell(row=r, column=c)) for r in (1, 3) for c in range(1, 9)))
wd = {k: v.width for k, v in w4.column_dimensions.items()}
check("sheet 4: widths A=28, B=40, C-H=42; freeze B4; autofilter A3:H29; no merges above the panel",
      [wd[get_column_letter(i)] for i in range(1, 9)] == [28, 40] + [42] * 6 and w4.freeze_panes == "B4"
      and w4.auto_filter.ref == "A3:H29" and not [m for m in w4.merged_cells.ranges if m.min_row < 31])
hdr_md, rows_md = BG.parse_groups()
GR = list(range(5, 11)) + list(range(12, 30))
BANDS = {4: "Comparison groups (2+ products)", 11: "Single-product functions (no comparator yet)"}
check("sheet 4: row map = header 3, band 4, C1-C6 rows 5-10, band 11, C7-C24 rows 12-29",
      all(plain(w4.cell(row=GR[i], column=1).value).startswith(f"C{i + 1} ") for i in range(24)))
bad = []
for br, t in BANDS.items():
    for c in range(1, 9):
        x = w4.cell(row=br, column=c)
        if not (plain(x.value) == (t if c == 1 else "") and x.font.name == "Arial" and x.font.sz == 10 and x.font.b
                and x.fill.fill_type == "solid" and str(x.fill.fgColor.rgb).endswith("D9D9D9")
                and x.border.left.style == "thin" and str(x.border.left.color.rgb).endswith("8EA9DB")):
            bad.append(x.coordinate)
check("sheet 4: band rows 4 and 11 (bold Arial 10, D9D9D9 across A:H, table borders, text in A only)", not bad, str(bad[:3]))
mism = sum(plain(w4.cell(row=GR[i], column=c).value) != v for i, row in enumerate(rows_md) for c, v in enumerate(row, 1))
md_lines = open(BG.MD, encoding="utf-8").read().splitlines()
tab = [l for l in md_lines[md_lines.index("## A. Group table"):md_lines.index("## Rewrite check")] if re.match(r"\| C\d", l)]
ind = 0
for i, l in enumerate(tab):
    cs = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", l.strip()[1:-1])]
    assert len(cs) == 8
    exp = [cs[0], "\n".join(x.strip() for x in cs[1].split("; "))] + ["\n".join(x.strip() for x in c.split("<br>")) for c in cs[2:]]
    ind += sum(plain(w4.cell(row=GR[i], column=c).value) != e for c, e in enumerate(exp, 1))
check("sheet 4: all 24 x 8 cells equal groups_v2.md (builder parse and independent parse; <br> -> newline)",
      mism == 0 and ind == 0 and len(tab) == 24, f"({mism}, {ind} mismatches)")
bf = []
for i in range(24):
    single = i >= 6
    k = i - 6 if single else i
    want = "595959" if single else None
    for c in range(1, 9):
        x = w4.cell(row=GR[i], column=c)
        col = _col(x.font.color)
        col = col[-6:] if isinstance(col, str) else col
        shaded = x.fill.fill_type == "solid" and str(x.fill.fgColor.rgb).endswith("DCE6F2")
        if (x.font.name, x.font.sz, bool(x.font.b), col) != ("Arial", 10, c == 1, want) or shaded != (k % 2 == 0) or \
                (not shaded and x.fill.fill_type is not None) or not (x.alignment.wrap_text and x.alignment.vertical == "top"
                and x.border.top.style == "thin" and str(x.border.top.color.rgb).endswith("8EA9DB")):
            bf.append(x.coordinate)
        if c >= 3:
            runs = list(x.value) if isinstance(x.value, CellRichText) else []
            if not runs:
                bf.append(x.coordinate + ":notrich")
                continue
            nref = 0
            for r_ in runs:
                t = r_.text if hasattr(r_, "text") else str(r_)
                f = r_.font if hasattr(r_, "font") else None
                ref = t.startswith("Refs: ")
                nref += ref
                rc = f.color.rgb[-6:] if f is not None and f.color is not None and isinstance(f.color.rgb, str) else None
                if f is None or (f.rFont, f.sz, rc) != (("Arial", 8, "808080") if ref else ("Arial", 10, want)):
                    bf.append(x.coordinate + ":run")
            if nref > 1:
                bf.append(x.coordinate + ":refs")
check("sheet 4: A bold; text runs Arial 10 (black C1-C6, 595959 C7-C24); 'Refs:' run Arial 8 808080; fills alternate "
      "per section; wrap/top/thin 8EA9DB", not bf, str(bf[:4]))
check("sheet 4: column B one function per line", all("; " not in plain(w4.cell(row=GR[i], column=2).value) and
      plain(w4.cell(row=GR[i], column=2).value).count("\n") == r[1].count("\n") for i, r in enumerate(rows_md)))
check("sheet 4: no template leftovers ('[Add additional', 'Example:')", not any(
    "[Add additional" in plain(c.value) or plain(c.value).startswith("Example:") for row in w4.iter_rows() for c in row))
hd = [ws3.cell(row=3, column=c).value for c in range(5, 34)]
P = 31
want4 = {(P, 1): "Coverage of sheet 3 functions", (P + 1, 1): "Sheet 3 column", (P + 1, 2): "Function (live link)",
         (P + 1, 3): "Groups listing it"}
for i in range(29):
    r = P + 2 + i
    L = get_column_letter(5 + i)
    want4[(r, 1)] = L
    want4[(r, 2)] = f"='{S3}'!{L}3"
    want4[(r, 3)] = f'=COUNTIF($B$5:$B$29,"*"&$B{r}&"*")'
f0, fN = P + 2, P + 30
rng_c = f"$C${f0}:$C${fN}"
tot_rows = [("Functions in no group", f'=COUNTIF({rng_c},0)'),
            ("Functions in more than one group", f'=COUNTIF({rng_c},">1")'),
            ("Groups including NeMo Guardrails", '=COUNTIF($B$5:$B$29,"*NeMo Guardrails:*")'),
            ("Groups including Llama Guard", '=COUNTIF($B$5:$B$29,"*Llama Guard:*")'),
            ("Groups including GovTech Sentinel", '=COUNTIF($B$5:$B$29,"*GovTech Sentinel:*")'),
            ("Groups including Amazon Bedrock", '=COUNTIF($B$5:$B$29,"*Amazon Bedrock*")'),
            ("Single-product rows", '=COUNTIF($H$5:$H$29,"*Single-product:*")')]
for k, (lab, f) in enumerate(tot_rows):
    want4[(fN + 1 + k, 1)] = lab
    want4[(fN + 1 + k, 3)] = f
got4 = {(c.row, c.column): plain(c.value) for row in w4.iter_rows(min_row=30) for c in row if c.value is not None}
check("sheet 4: panel (%d cells) below the table holds exactly the specified formulas; title row 31; functions rows %d-%d; totals %d-%d"
      % (len(want4), f0, fN, fN + 1, fN + 7), got4 == want4,
      "" if got4 == want4 else str([(k, got4.get(k), want4.get(k)) for k in sorted(set(got4) | set(want4))
                                    if got4.get(k) != want4.get(k)][:4]))
check("sheet 4: row 30 empty; panel title merged A:C", all(c.value is None for c in w4[30]) and
      "A31:C31" in {str(m) for m in w4.merged_cells.ranges})
check("sheet 4: guards - headers E3..AG3 no substring of another, <=253 chars, no * ? ~",
      all(not any(ch in h for ch in "*?~") and len(h) + 2 <= 255 for h in hd) and
      not any(h != o and h.lower() in o.lower() for h in hd for o in hd))
B = [plain(w4.cell(row=r, column=2).value).lower() for r in range(5, 30)]
cnt = [sum(h.lower() in b for b in B) for h in hd]
mt = {"no group": cnt.count(0), "more than one group": sum(c > 1 for c in cnt)}
for lab, key in (("NeMo Guardrails", "nemo guardrails:"), ("Llama Guard", "llama guard:"),
                 ("GovTech Sentinel", "govtech sentinel:"), ("Amazon Bedrock", "amazon bedrock")):
    mt[lab] = sum(key in b for b in B)
mt["single-product rows"] = sum("single-product:" in plain(w4.cell(row=r, column=8).value).lower() for r in range(5, 30))
print("   Python mirror, groups listing each function:", {get_column_letter(5 + i): c for i, c in enumerate(cnt)})
print("   Python mirror totals:", mt)
print("   functions in more than one group:", [get_column_letter(5 + i) for i, c in enumerate(cnt) if c > 1])
check("sheet 4: Python mirror - 0 functions in no group", mt["no group"] == 0)
check("sheet 4: Python mirror - >1 group 7, NeMo 17, Llama Guard 7, Sentinel 8, Bedrock 2, single-product 18",
      (mt["more than one group"], mt["NeMo Guardrails"], mt["Llama Guard"], mt["GovTech Sentinel"], mt["Amazon Bedrock"],
       mt["single-product rows"]) == (7, 17, 7, 8, 2, 18))
mk = [c.coordinate for row in w4.iter_rows() for c in row if "**" in plain(c.value) or "`" in plain(c.value)]
check("sheet 4: no ** or backticks", not mk)
nx, bx = xml_non_arial(XLSX, S4)
nonar = [c.coordinate for row in w4.iter_rows() for c in row
         if type(c).__name__ != "MergedCell" and c.value is not None and c.font.name != "Arial"]
check("sheet 4: all Arial (cells and raw XML)", not nonar and bx == 0, f"(cells {len(nonar)}, raw {bx}/{nx})")

# ---------------------------------------------------------------- XML hygiene
ssp, emp, norpr, n, wf = BI.hygiene(XLSX)
check("XML hygiene (all parts well-formed; worksheets + sharedStrings)", ssp == 0 and emp == 0 and norpr == 0 and wf,
      f"(unpreserved <t>={ssp}, <r> without rPr={norpr}, empty <t>={emp}, parts checked={n}, well-formed={wf})")

print("\nSUMMARY: %d/%d checks passed" % (sum(RESULTS), len(RESULTS)))
sys.exit(0 if all(RESULTS) else 1)
